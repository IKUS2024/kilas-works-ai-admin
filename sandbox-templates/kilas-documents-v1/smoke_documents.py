"""Fixed build-time document smoke. No model source, network or customer data."""
import io
import tempfile
from pathlib import Path
import zipfile


def smoke():
    from reportlab.pdfgen import canvas
    from docx import Document
    from openpyxl import Workbook,load_workbook
    from openpyxl.chart import BarChart,Reference
    from pptx import Presentation
    output={}
    raw=io.BytesIO();pdf=canvas.Canvas(raw);pdf.drawString(40,800,'Synthetic revenue 410000 cost 226000 profit 184000');pdf.save();output['report.pdf']=raw.getvalue()
    doc=Document();doc.add_heading('Synthetic sales',0);doc.add_paragraph('Revenue 410000, cost 226000, profit 184000.');raw=io.BytesIO();doc.save(raw);output['report.docx']=raw.getvalue()
    book=Workbook();sheet=book.active
    for row in [('Item','Revenue','Cost','Profit'),('Kopi',200000,120000,'=B2-C2'),('Teh',120000,56000,'=B3-C3'),('Roti',90000,50000,'=B4-C4'),('Total','=SUM(B2:B4)','=SUM(C2:C4)','=SUM(D2:D4)')]:sheet.append(row)
    chart=BarChart();chart.add_data(Reference(sheet,min_col=2,max_col=4,min_row=1,max_row=4),titles_from_data=True);sheet.add_chart(chart,'F2');raw=io.BytesIO();book.save(raw);output['report.xlsx']=raw.getvalue()
    deck=Presentation();slide=deck.slides.add_slide(deck.slide_layouts[5]);slide.shapes.title.text='Synthetic profit 184000';raw=io.BytesIO();deck.save(raw);output['report.pptx']=raw.getvalue()
    assert output['report.pdf'].startswith(b'%PDF-') and b'%%EOF' in output['report.pdf'][-1024:]
    assert Document(io.BytesIO(output['report.docx'])).paragraphs[1].text.endswith('184000.')
    checked=load_workbook(io.BytesIO(output['report.xlsx']));assert checked.active['D5'].value=='=SUM(D2:D4)' and len(checked.active._charts)==1;checked.close()
    assert Presentation(io.BytesIO(output['report.pptx'])).slides[0].shapes.title.text=='Synthetic profit 184000'
    for name,data in output.items():
        if name.endswith(('docx','xlsx','pptx')):
            with zipfile.ZipFile(io.BytesIO(data)) as archive:assert archive.testzip() is None
    return output

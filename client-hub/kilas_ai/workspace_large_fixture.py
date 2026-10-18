"""Fixed synthetic large-document guest fixture. Never execute on web host."""
import importlib.util
import io
import json
from pathlib import Path
import zipfile

PACKAGES=('reportlab','docx','openpyxl','pptx')
SIZES={'report.pdf':8388608,'report.docx':2097152,'report.xlsx':3145728,'report.pptx':3145728}
TOTALS={'revenue':410000,'cost':226000,'profit':184000}


def pad_office(raw,size):
    def build(length):
        output=io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(raw)) as source,zipfile.ZipFile(output,'w') as target:
            for item in source.infolist():target.writestr(item,source.read(item))
            target.writestr('customXml/kilas-capacity.xml',b'<capacity>'+b'x'*length+b'</capacity>',compress_type=zipfile.ZIP_STORED)
        return output.getvalue()
    empty=build(0);result=build(size-len(empty));assert len(result)==size
    return result


def documents():
    # openpyxl supports pure Python; avoid optional NumPy/OpenBLAS virtual-memory
    # reservations under the unchanged 256MiB address-space cap.
    import sys
    sys.modules['numpy']=None
    from reportlab.pdfgen import canvas
    from docx import Document
    from openpyxl import Workbook
    from openpyxl.chart import BarChart,Reference
    from pptx import Presentation
    from pptx.util import Inches
    summary='Synthetic sales: revenue 410000; cost 226000; profit 184000.'
    def build_pdf(padding):
        stream=io.BytesIO();pdf=canvas.Canvas(stream,pageCompression=0);pdf.drawString(40,800,summary)
        pdf._code.append('%'+('x'*padding));pdf.save();return stream.getvalue()
    padding=0;raw=build_pdf(padding)
    for _ in range(4):
        if len(raw)==SIZES['report.pdf']:break
        padding+=SIZES['report.pdf']-len(raw);raw=build_pdf(padding)
    assert len(raw)==SIZES['report.pdf']
    files={'report.pdf':raw}
    document=Document();document.add_heading('Synthetic sales',0);document.add_paragraph(summary)
    table=document.add_table(rows=1,cols=4)
    for cell,text in zip(table.rows[0].cells,('Item','Units','Unit price','Unit cost')):cell.text=text
    for values in (('Kopi',10,20000,12000),('Teh',8,15000,7000),('Roti',5,18000,10000)):
        for cell,value in zip(table.add_row().cells,values):cell.text=str(value)
    stream=io.BytesIO();document.save(stream);files['report.docx']=pad_office(stream.getvalue(),SIZES['report.docx'])
    book=Workbook();sheet=book.active
    for row in [('Item','Units','Price','Unit cost','Revenue','Cost','Profit'),('Kopi',10,20000,12000,'=B2*C2','=B2*D2','=E2-F2'),('Teh',8,15000,7000,'=B3*C3','=B3*D3','=E3-F3'),('Roti',5,18000,10000,'=B4*C4','=B4*D4','=E4-F4'),('Total','','','','=SUM(E2:E4)','=SUM(F2:F4)','=SUM(G2:G4)')]:sheet.append(row)
    chart=BarChart();chart.add_data(Reference(sheet,min_col=5,max_col=7,min_row=1,max_row=4),titles_from_data=True);sheet.add_chart(chart,'I2')
    stream=io.BytesIO();book.save(stream);files['report.xlsx']=pad_office(stream.getvalue(),SIZES['report.xlsx'])
    deck=Presentation();slide=deck.slides.add_slide(deck.slide_layouts[5]);slide.shapes.title.text='Synthetic sales'
    box=slide.shapes.add_textbox(Inches(1),Inches(2),Inches(8),Inches(2));box.text_frame.text=summary
    stream=io.BytesIO();deck.save(stream);files['report.pptx']=pad_office(stream.getvalue(),SIZES['report.pptx'])
    assert {name:len(raw) for name,raw in files.items()}==SIZES
    return files


def main():
    import boundary_checks as boundary
    inventory={name:importlib.util.find_spec(name) is not None for name in PACKAGES}
    report={'packages':inventory,'totals':TOTALS,'boundary':None,'documents':False}
    stage='boundary'
    try:
        boundary.main()  # Includes real EFBIG, data/inode ENOSPC, CPU/memory/FD/syscall checks.
        report['boundary']=json.loads(Path('output/proof.json').read_text())
        Path('output/proof.json').unlink();Path('output/answer.txt').unlink()
        stage='packages'
        if not all(inventory.values()):raise RuntimeError('required_packages_missing')
        stage='documents'
        for name,raw in documents().items():Path('output',name).write_bytes(raw)
        report['documents']=True;report['sizes']=SIZES
        print('KILAS_LARGE_OK '+json.dumps(report,separators=(',',':')))
    except BaseException:
        report['failed_stage']=stage
        print('KILAS_LARGE_FAILED '+json.dumps(report,separators=(',',':')))
        raise SystemExit(1)


if __name__=='__main__':main()

"""One reviewed, bounded VM acceptance for root-owned offline document wheels.

No template build, public execution endpoint, credential change or model call.
The fixed release operation is consumed durably before dispatch, including on
failure. General access remains subject to the existing entitlement and meters.
"""
import hashlib
import io
import json
import logging
import os
from pathlib import Path
import re
import secrets
import threading
from datetime import datetime, timedelta, timezone
import db
from . import autonomous_store as store, usage

KEY = 'document-runtime-20261019-v1'
RELEASE_ENABLED = True  # Reviewed one-shot release; expiry and durable key still apply.
OWNER_ID = 9
SERVICE_ID = 'srv-da7ti2psrm7s73dh9i2g'
EXPIRES_AT = datetime(2026, 10, 20, 10, 2, 37, tzinfo=timezone.utc)
TASK_BUDGET_MICRO_USD = 1000000
VM_RESERVE_MICRO_USD = 8880  # 60 s at the entire Hobby 8-vCPU/8-GiB ceiling.
PROBE_COMMAND = "/usr/bin/python3 -I -c 'import json,platform,sys,sysconfig; print(json.dumps({\"abi\":\"cp\"+str(sys.version_info.major)+str(sys.version_info.minor),\"python\":platform.python_version(),\"machine\":platform.machine(),\"libc\":platform.libc_ver(),\"free_threaded\":bool(sysconfig.get_config_var(\"Py_GIL_DISABLED\"))}))'"


def fingerprint():
    from . import workspace_large_outputs as large
    root = Path(__file__).parent
    names = ('workspace_document_runtime.py', 'workspace_document_guest.py',
             'workspace_document_fixture.py', 'workspace_document_bundle.py',
             'workspace_document_bundle.lock.json')
    return hashlib.sha256(large.fingerprint().encode()+b''.join((root/name).read_bytes() for name in names)).hexdigest()


def approval_valid():
    # This narrowly scoped release approval expires; a failed/consumed operation
    # can never become a second VM by restart, key rotation or source changes.
    from . import workspace_execution as base
    return bool(RELEASE_ENABLED and base.MAX_SECONDS == 60 and base.RESERVE_MICRO_USD == VM_RESERVE_MICRO_USD
                and os.environ.get('RENDER_SERVICE_ID') == SERVICE_ID and
                os.environ.get('E2B_API_KEY') and store.now() < EXPIRES_AT)


def fixture():
    root = Path(__file__).parent
    boundary = (root/'workspace_acceptance_fixture.py').read_text()
    needle = "target.write(b'x'*65536)"
    if boundary.count(needle) != 1:
        raise ValueError('document_fixture_baseline_changed')
    return {'main.py':(root/'workspace_document_fixture.py').read_text(),
            'boundary_checks.py':boundary.replace(needle, "target.write(b'x'*8388608)")}


def receipt():
    row = db.query_one("SELECT e.outcome_json FROM kilas_workspace_executions e JOIN kilas_agent_jobs j ON j.id=e.job_id WHERE e.operation_key=? AND e.state='SUCCEEDED' AND e.cleanup_confirmed=1 AND j.status='COMPLETED'", ('execution:'+KEY,))
    value = json.loads(row['outcome_json']) if row else {}
    if (value.get('document_proof_passed') is True and value.get('download_checked') is True
            and value.get('document_fingerprint') == fingerprint() and valid_runtime(value.get('runtime_evidence'))
            and isinstance(value.get('template_id'),str) and re.fullmatch('[a-z0-9-]{3,64}',value['template_id'])
            and value['template_id'] != 'base'):
        return value
    return None


def accepted():
    return receipt() is not None


def authorized_for(job, step, files, entrypoint):
    return bool(approval_valid() and job['user_id'] == OWNER_ID and
                step['idempotency_key'] == KEY and entrypoint == 'main.py' and files == fixture())


def probe(sandbox, expected=None):
    from . import workspace_document_bundle as bundle
    command = sandbox.commands.run(PROBE_COMMAND, user='root', timeout=3, request_timeout=5)
    if command.exit_code != 0 or command.stderr or len(command.stdout) > 512:
        raise ValueError('document_runtime_unavailable')
    try:
        value = json.loads(command.stdout)
        if set(value) != {'abi','python','machine','libc','free_threaded'} or value['machine'] != 'x86_64' or value['free_threaded'] is not False:
            raise ValueError()
        if not re.fullmatch(r'3\.(?:10|11|12|13|14)\.\d+', value['python']):
            raise ValueError()
        if value['libc'][0] != 'glibc' or tuple(map(int, value['libc'][1].split('.')[:2])) < (2, 28):
            raise ValueError()
        spec = bundle.describe(value['abi'])
        if expected and value['python'] != expected['runtime_evidence']['python']:
            raise ValueError()
        return spec
    except (ValueError, TypeError, KeyError, IndexError):
        raise ValueError('document_runtime_unavailable') from None


def valid_runtime(value):
    from . import workspace_document_bundle as bundle
    if not isinstance(value, dict) or set(value) != {'abi','python','manifest_sha256','versions','root_owned','network_install','host_execution'}:
        return False
    try:
        expected = {w['project'].lower().replace('_','-'):w['version'] for w in bundle.describe(value['abi'])['wheels']}
        return bool(value['manifest_sha256'] == bundle.manifest_digest() and value['versions'] == expected
                    and re.fullmatch(r'3\.(?:10|11|12|13|14)\.\d+', value['python'])
                    and value['abi'] == 'cp'+''.join(value['python'].split('.')[:2])
                    and value['root_owned'] is True and value['network_install'] is False and value['host_execution'] is False)
    except (ValueError, TypeError, KeyError):
        return False


def passed(outcome):
    from . import workspace_execution as base, workspace_acceptance_fixture as boundary
    from . import workspace_document_fixture as fixed, workspace_large_outputs as large, workspace_output_formats as formats
    if not (type(outcome.get('exit_code')) is int and outcome['exit_code'] == 0 and outcome.get('timed_out') is False
            and outcome.get('customer_metered') is False and outcome.get('cleanup_confirmed') is True
            and outcome.get('sandbox_created') is True and outcome.get('stderr') == ''
            and type(outcome.get('elapsed_ms')) is int and 0 <= outcome['elapsed_ms'] < 60000
            and valid_runtime(outcome.get('runtime_evidence'))):
        return False
    try:
        base.validated_evidence(outcome['boundary_evidence'], True)
        lines = outcome['stdout'].splitlines()
        if len(lines) != 2 or lines[0] != 'KILAS_ACCEPTANCE_OK' or not lines[1].startswith('KILAS_DOCUMENT_OK '):
            return False
        value = json.loads(lines[1].split(' ', 1)[1])
        if set(value) != {'boundary','sizes','pages','slides'} or value['pages'] != fixed.PAGES or value['slides'] != fixed.SLIDES:
            return False
        report = value['boundary']
        if report['answer'] != 42 or report['checks'] != boundary.CHECKS:
            return False
        data, inodes = report['evidence']['data'], report['evidence']['inodes']
        if not (0 < data['written_bytes'] <= data['available_before'] <= boundary.BYTES
                and data['available_before'] < data['requested_bytes'] <= boundary.BYTES+boundary.CHUNK and data['errno'] == 28
                and 0 < inodes['created_files'] == inodes['available_before'] <= boundary.INODES and inodes['errno'] == 28):
            return False
        if outcome['collector_evidence'] != {'checks':['file_limit','job_limit','file_count'],'private_staging':True,'post_stop':True}:
            return False
        files = outcome['large_files']
        if len(files) != 4 or {f['filename'] for f in files} != {'small.pdf','large.pdf','small.pptx','large.pptx'}:
            return False
        if value['sizes'] != {f['filename']:len(f['content']) for f in files} or sum(value['sizes'].values()) > large.MAX_JOB:
            return False
        for file in files:
            small = file['filename'].startswith('small.')
            if not (0 < len(file['content']) <= 65536 if small else 65536 < len(file['content']) <= large.MAX_FILE):
                return False
            formats.validate(file['filename'], file['mime_type'], file['content'], large=True)
            if file['filename'].endswith('.pdf'):
                from pypdf import PdfReader
                pages = PdfReader(io.BytesIO(file['content']), strict=True).pages
                if len(pages) != (1 if small else fixed.PAGES) or not all('184000' in page.extract_text() for page in pages):
                    return False
            else:
                from pptx import Presentation
                slides = Presentation(io.BytesIO(file['content'])).slides
                if len(slides) != (1 if small else fixed.SLIDES) or not all('184000' in slide.shapes.title.text for slide in slides):
                    return False
        for label in ('host_memory_before','host_memory_after'):
            memory = outcome[label]
            if not 0 < memory['current'] <= memory['peak'] <= 448*1024*1024 or memory['limit'] > 512*1024*1024:
                return False
        return True
    except (ValueError, TypeError, KeyError, AssertionError):
        return False


def claim():
    # The unique step key and job insert are one transaction. A racing worker
    # rolls back its entire job; a process crash cannot recreate a consumed key.
    if db.query_one('SELECT id FROM kilas_agent_steps WHERE idempotency_key=?', (KEY,)):
        return None
    token = secrets.token_hex(24)
    try:
        with store.transaction() as conn:
            if db.BACKEND == 'postgres':
                usage._query(conn, 'SELECT pg_advisory_xact_lock(74110601)', ())
            if usage._query(conn, 'SELECT id FROM kilas_agent_steps WHERE idempotency_key=?', (KEY,), one=True):
                return None
            spent = usage._query(conn, "SELECT COALESCE(SUM(reserved_micro_usd),0) FROM kilas_workspace_executions WHERE operation_key LIKE 'execution:document-runtime-20261019-%'", (), one=True)[0]
            if spent+VM_RESERVE_MICRO_USD > TASK_BUDGET_MICRO_USD:
                raise ValueError('document_release_budget_exhausted')
            count = usage._query(conn, "SELECT COUNT(*) FROM kilas_agent_jobs WHERE user_id=? AND status NOT IN ('COMPLETED','FAILED','STOPPED')", (OWNER_ID,), one=True)[0]
            if count >= 20:
                raise ValueError('active_job_limit')
            ident = usage._query(conn, "INSERT INTO kilas_agent_jobs(user_id,title,instruction,mode,status,lease_token,lease_until,next_wake_at,max_steps,max_attempts) VALUES (?,?,?,'ONE_SHOT','RUNNING',?,?,NULL,1,1) RETURNING id", (OWNER_ID,'Document runtime verification','Fixed synthetic PDF and PPTX acceptance',token,store.stamp(store.now()+timedelta(seconds=120))), one=True)[0]
            step_id = usage._query(conn, "INSERT INTO kilas_agent_steps(job_id,sequence,worker,action,instruction,input_json,idempotency_key,status,attempts) VALUES (?,1,'EXECUTION','python','Fixed document runtime proof','{}',?,'RUNNING',1) RETURNING id", (ident,KEY), one=True)[0]
        return store.get(OWNER_ID,ident), dict(db.query_one('SELECT * FROM kilas_agent_steps WHERE id=?',(step_id,)))
    except Exception:
        if db.query_one('SELECT id FROM kilas_agent_steps WHERE idempotency_key=?', (KEY,)):
            return None
        raise


def run(app):
    from . import workspace_execution as base, workspace_large_execution as execution, workspace_large_outputs as large
    from . import workspace_document_bundle as bundle, workspace_execution_worker as worker, autonomous_runner as runner
    from .workspace_large_proof import downloaded
    from .agent_workers import Result
    with app.app_context():
        if not approval_valid() or db.query_one('SELECT id FROM kilas_agent_steps WHERE idempotency_key=?', (KEY,)):
            return
        phase = 'capability'
        try:
            if not base.capability(OWNER_ID)['available']:
                raise ValueError('existing_execution_unavailable')
            phase = 'host_memory'
            large.memory_admission()
            phase = 'wheel_cache'
            bundle.prepare()  # Free verified downloads; always BEFORE any VM.
            phase = 'durable_claim'
            claimed = claim()
        except Exception as error:
            reason = str(error) if isinstance(error, bundle.BundleError) else str(error) if str(error) in {
                'existing_execution_unavailable', 'large_host_memory_unverified', 'large_host_memory_busy',
                'document_release_budget_exhausted', 'active_job_limit'} else 'preflight_error'
            memory = {}
            if phase == 'host_memory':
                for name in ('memory.max', 'memory.current', 'memory.peak'):
                    try:
                        value = (Path('/sys/fs/cgroup') / name).read_text().strip()
                        memory[name] = int(value) if value.isdecimal() else 'unbounded' if value == 'max' else 'unavailable'
                    except (OSError, ValueError):
                        memory[name] = 'unavailable'
            logging.getLogger(__name__).warning('KILAS_DOCUMENT_PROOF state=PREFLIGHT_BLOCKED phase=%s reason=%s memory=%s provider_dispatched=False',
                phase, reason, json.dumps(memory, sort_keys=True))
            return
        if not claimed:
            return
        job, step = claimed
        try:
            outcome = execution.execute(job, step, fixture(), _document_proof=True)
            verified = passed(outcome)
            rendered = worker.result(outcome)
            result = rendered if verified else Result('FAILED','Document verification failed.',{'reason':'document_proof_failed'},rendered.artifacts[:1])
            finished = runner.finish(job,job['lease_token'],step,result)
            checked = bool(verified and finished and outcome.get('document_fingerprint') == fingerprint()
                           and downloaded(app,job,outcome['large_files']))
            memory = large.host_memory()
            checked = bool(checked and memory['peak'] <= 448*1024*1024 and memory['current'] <= memory['peak'])
            public = json.loads(db.query_one('SELECT outcome_json FROM kilas_workspace_executions WHERE id=?',(outcome['run_id'],))['outcome_json'])
            public.update(document_proof_passed=checked,download_checked=checked,host_memory_final=memory)
            base.update(outcome['run_id'],outcome_json=json.dumps(public))
            logging.getLogger(__name__).warning('KILAS_DOCUMENT_PROOF state=%s run_id=%s cleanup=%s bytes=%s compute_micro=%s',
                'PASSED' if checked else 'FAILED',outcome['run_id'],outcome['cleanup_confirmed'],
                sum(len(f['content']) for f in outcome.get('large_files',[])),outcome['compute_cost_upper_estimate_micro_usd'])
        except Exception:
            runner.finish(job,job['lease_token'],step,Result('FAILED','Document verification failed.',{'reason':'document_proof_failed'}))
            logging.getLogger(__name__).warning('KILAS_DOCUMENT_PROOF state=FAILED retry=False')


def maybe_start(app):
    # Internal fixed release task only. No endpoint accepts user code or grants
    # general availability; a complete durable acceptance receipt grants it.
    logging.getLogger(__name__).warning('KILAS_DOCUMENT_GATE requested=%s accepted=%s template_builds=0 vm_timeout_seconds=60 release_budget_micro_usd=%s reserve_per_attempt_micro_usd=%s host_execution=False',
        approval_valid(),accepted(),TASK_BUDGET_MICRO_USD,VM_RESERVE_MICRO_USD)
    if approval_valid() and not app.testing:
        threading.Thread(target=run,args=(app,),name='document-runtime-proof',daemon=True).start()

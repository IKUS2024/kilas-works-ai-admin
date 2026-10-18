"""One explicitly approved large-profile VM; no automatic replacement/retry."""
import hashlib,json,logging,os,secrets,threading
from datetime import timedelta
from pathlib import Path
import db
from . import autonomous_store as store,usage

KEY='large-document-acceptance-v1'
POLICY='e2b-large-8m-16m-60s-v1'
OWNER_ID=9


def fixture():
    root=Path(__file__).parent
    boundary=(root/'workspace_acceptance_fixture.py').read_text()
    needle="target.write(b'x'*65536)"
    if boundary.count(needle)!=1:raise ValueError('large_fixture_baseline_changed')
    boundary=boundary.replace(needle,"target.write(b'x'*8388608)")
    return {'main.py':(root/'workspace_large_fixture.py').read_text(),'boundary_checks.py':boundary}


def authorized():
    return os.environ.get('KILAS_LARGE_PROOF_APPROVAL')==POLICY and bool(os.environ.get('E2B_API_KEY'))


def accepted():
    from . import workspace_large_outputs as large
    row=db.query_one("SELECT e.outcome_json FROM kilas_workspace_executions e JOIN kilas_agent_jobs j ON j.id=e.job_id WHERE e.operation_key=? AND e.state='SUCCEEDED' AND e.cleanup_confirmed=1 AND j.status='COMPLETED'",('execution:'+KEY,))
    value=json.loads(row['outcome_json']) if row else {}
    return bool(value.get('large_proof_passed') is True and value.get('download_checked') is True and value.get('large_proof_fingerprint')==large.fingerprint())


def passed(outcome):
    from . import workspace_execution as base,workspace_acceptance_fixture as boundary,workspace_large_fixture as fixture
    from . import workspace_large_outputs as large
    if not (outcome.get('exit_code')==0 and type(outcome.get('exit_code')) is int and outcome.get('timed_out') is False and outcome.get('customer_metered') is False and outcome.get('cleanup_confirmed') is True and outcome.get('sandbox_created') is True and outcome.get('stderr')=='' and type(outcome.get('elapsed_ms')) is int and 0<=outcome['elapsed_ms']<60000):return False
    try:
        base.validated_evidence(outcome['boundary_evidence'],True)
        lines=outcome['stdout'].splitlines()
        if len(lines)!=2 or lines[0]!='KILAS_ACCEPTANCE_OK' or not lines[1].startswith('KILAS_LARGE_OK '):return False
        value=json.loads(lines[1].split(' ',1)[1]);report=value['boundary']
        if set(value)!={'packages','totals','boundary','documents','sizes'} or value['packages']!={k:True for k in fixture.PACKAGES} or value['totals']!=fixture.TOTALS or value['documents'] is not True or value['sizes']!=fixture.SIZES:return False
        if report['answer']!=42 or report['checks']!=boundary.CHECKS:return False
        data,inodes=report['evidence']['data'],report['evidence']['inodes']
        if not (0<data['written_bytes']<=data['available_before']<=boundary.BYTES and data['available_before']<data['requested_bytes']<=boundary.BYTES+boundary.CHUNK and data['errno']==28 and 0<inodes['created_files']==inodes['available_before']<=boundary.INODES and inodes['errno']==28):return False
        if outcome['collector_evidence']!={'checks':['file_limit','job_limit','file_count'],'private_staging':True,'post_stop':True}:return False
        files=outcome['large_files']
        if len(files)!=4 or {f['filename']:len(f['content']) for f in files}!=fixture.SIZES or sum(len(f['content']) for f in files)!=large.MAX_JOB:return False
        for label in ('host_memory_before','host_memory_after'):
            memory=outcome[label]
            if not 0<memory['current']<=memory['peak']<=448*1024*1024 or memory['limit']>512*1024*1024:return False
        validate_documents(files)
        return True
    except (ValueError,TypeError,KeyError,AssertionError):return False


def validate_documents(files):
    import gc,io
    from . import workspace_output_formats as formats
    for file in files:
        formats.validate(file['filename'],file['mime_type'],file['content'],large=True)
        stream=io.BytesIO(file['content'])
        if file['filename']=='report.pdf':
            from pypdf import PdfReader
            value=PdfReader(stream)
            if len(value.pages)!=1 or '410000' not in value.pages[0].extract_text():raise ValueError('large_pdf_content')
        elif file['filename']=='report.docx':
            from docx import Document
            value=Document(stream)
            if len(value.tables)!=1 or '184000' not in ' '.join(p.text for p in value.paragraphs):raise ValueError('large_docx_content')
        elif file['filename']=='report.xlsx':
            from openpyxl import load_workbook
            value=load_workbook(stream)
            if value.active['G2'].value!='=E2-F2' or value.active['G5'].value!='=SUM(G2:G4)' or len(value.active._charts)!=1:raise ValueError('large_xlsx_content')
            value.close()
        elif file['filename']=='report.pptx':
            from pptx import Presentation
            value=Presentation(stream)
            if len(value.slides)!=1 or not any('184000' in shape.text for shape in value.slides[0].shapes if shape.has_text_frame):raise ValueError('large_pptx_content')
        del value;stream.close();gc.collect()


def downloaded(app,job,files):
    # Real owner-scoped handler/response bytes under synthetic request contexts.
    # No login is created; external authentication/browser E2E is not claimed.
    from flask import session
    from werkzeug.exceptions import NotFound
    from . import work_artifacts
    from .autonomous_routes import artifact
    rows=work_artifacts.listing(OWNER_ID,job_id=job['id'])
    expected={f['filename']:hashlib.sha256(f['content']).hexdigest() for f in files}
    for row in rows:
        if row['name'] not in expected:continue
        path=f'/kilas-ai/agent/jobs/{job["id"]}/artifacts/{row["id"]}?download=1'
        with app.test_request_context(path):
            session['user_id']=OWNER_ID;response=artifact(job['id'],row['id']);response.direct_passthrough=False
            try:
                if response.status_code!=200 or hashlib.sha256(response.get_data()).hexdigest()!=expected.pop(row['name']) or 'attachment' not in response.headers.get('Content-Disposition','') or response.headers.get('X-Content-Type-Options')!='nosniff':return False
            finally:response.close()
        with app.test_request_context(path):
            session['user_id']=0
            try:artifact(job['id'],row['id'])
            except NotFound:pass
            else:return False
    return not expected


def run(app):
    from . import workspace_execution as base,workspace_large_execution as execution,workspace_large_outputs as large
    from . import workspace_execution_worker as worker,autonomous_runner as runner
    from .agent_workers import Result
    with app.app_context():
        if not authorized() or db.query_one('SELECT id FROM kilas_workspace_executions WHERE operation_key=?',('execution:'+KEY,)):return
        try:large.memory_admission()
        except ValueError as error:
            logging.getLogger(__name__).warning('KILAS_LARGE_PROOF preflight=%s',str(error));return
        ident=store.create(OWNER_ID,'Verify large isolated document outputs',conversation_id=None,wake_at=store.now()+timedelta(days=3650));token=secrets.token_hex(24)
        db.execute("UPDATE kilas_agent_jobs SET status='RUNNING',lease_token=?,lease_until=?,next_wake_at=NULL WHERE id=?",(token,store.stamp(store.now()+timedelta(seconds=120)),ident))
        with store.transaction() as conn:
            step_id=usage._query(conn,"INSERT INTO kilas_agent_steps(job_id,sequence,worker,action,instruction,input_json,idempotency_key,status,attempts) VALUES (?,1,'EXECUTION','python','Fixed large document proof','{}',?,'RUNNING',1) RETURNING id",(ident,KEY),one=True)[0]
        job=store.get(OWNER_ID,ident);step=dict(db.query_one('SELECT * FROM kilas_agent_steps WHERE id=?',(step_id,)))
        try:
            outcome=execution.execute(job,step,fixture(),_proof=True)
            verified=passed(outcome)
            result=worker.result(outcome)
            if not verified:result=Result('FAILED','Large sandbox verification failed.',{'reason':'large_proof_failed'},result.artifacts,verified=False)
            runner.finish(job,token,step,result)
            checked=bool(verified and outcome.get('large_proof_fingerprint')==large.fingerprint() and downloaded(app,job,outcome['large_files']))
            memory=large.host_memory();checked=bool(checked and memory['peak']<=448*1024*1024 and memory['current']<=memory['peak'])
            public=json.loads(db.query_one('SELECT outcome_json FROM kilas_workspace_executions WHERE id=?',(outcome['run_id'],))['outcome_json'])
            public.update(large_proof_passed=checked,download_checked=checked,host_memory_final=memory)
            base.update(outcome['run_id'],outcome_json=json.dumps(public))
            logging.getLogger(__name__).warning('KILAS_LARGE_PROOF state=%s run_id=%s cleanup=%s bytes=%s compute_micro=%s','PASSED' if checked else 'FAILED',outcome['run_id'],outcome['cleanup_confirmed'],sum(len(f['content']) for f in outcome.get('large_files',[])),outcome['compute_cost_upper_estimate_micro_usd'])
        except Exception:
            runner.finish(job,token,step,Result('FAILED','Large sandbox verification failed.',{'reason':'large_proof_failed'}))
            logging.getLogger(__name__).warning('KILAS_LARGE_PROOF state=FAILED')


def maybe_start(app):
    from . import workspace_large_outputs as large,workspace_execution as base
    candidate=next((row['id'] for row in db.query_all('SELECT id FROM users WHERE id<>? ORDER BY id',(OWNER_ID,)) if base.capability(row['id'])['available']),None)
    available=bool(candidate is not None and large.capability(candidate)['available'])
    logging.getLogger(__name__).warning('KILAS_LARGE_GATE proof_requested=%s requested=%s accepted=%s nonowner_entitled_available=%s customer_reservation_required=True fingerprint=%s max_file=%s max_job=%s storage_ceiling=%s concurrency=1 host_execution=False',authorized(),os.environ.get('KILAS_LARGE_OUTPUT_APPROVAL')==POLICY,accepted(),available,large.fingerprint(),large.MAX_FILE,large.MAX_JOB,large.STORAGE_CEILING)
    if authorized():threading.Thread(target=run,args=(app,),name='fixed-large-proof',daemon=True).start()

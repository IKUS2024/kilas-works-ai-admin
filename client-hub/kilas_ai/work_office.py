"""Safe local Office formatting: no macros, remote templates or untrusted formulas."""
import csv
import io
import json
import math
import re
from . import work_artifacts


def checked_formula(value):
    """Small arithmetic/range subset; no external workbooks, DDE, URLs or macros."""
    if not isinstance(value,str) or len(value)>200 or not re.fullmatch(r'=[A-Za-z0-9:,+*/(). <>=-]+',value):raise ValueError('unsafe_spreadsheet_formula')
    functions=re.findall(r'([A-Za-z]+)\s*\(',value)
    if any(name.upper() not in ('SUM','AVERAGE','MIN','MAX','COUNT','COUNTA','ROUND','IF') for name in functions):raise ValueError('unsafe_spreadsheet_function')
    for col,row in re.findall(r'([A-Z]+)(\d+)',value.upper()):
        if len(col)>1 or not 1<=int(row)<=1000:raise ValueError('formula_range_limit')
    # Reject arbitrary named ranges; cell references and allowlisted functions only.
    residue=re.sub(r'[A-Za-z]+\d+|(?:SUM|AVERAGE|MIN|MAX|COUNT|COUNTA|ROUND|IF)\s*(?=\()', '',value,flags=re.I)
    if re.search(r'[A-Za-z_]',residue) or value.count('(')!=value.count(')'):raise ValueError('invalid_spreadsheet_formula')
    return value


def spreadsheet(source,output):
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill
    from openpyxl.chart import BarChart,LineChart,PieChart,Reference
    from openpyxl.workbook.properties import CalcProperties
    # CSV remains supported for existing clients; formulas require explicit JSON.
    structured=source.lstrip().startswith('{')
    spec=json.loads(source) if structured else {'sheets':[{'name':'Data','rows':list(csv.reader(io.StringIO(source)))}]}
    sheets=spec.get('sheets')
    if not isinstance(sheets,list) or not 1<=len(sheets)<=8:raise ValueError('invalid_workbook')
    book=Workbook();book.remove(book.active);names=set()
    for item in sheets:
        name=item.get('name','Data');rows=item.get('rows')
        if not isinstance(name,str) or not 1<=len(name)<=31 or re.search(r'[\\/*?:\[\]]',name) or name.casefold() in names:raise ValueError('invalid_sheet_name')
        if not isinstance(rows,list) or not 2<=len(rows)<=1000 or not isinstance(rows[0],list) or not 1<=len(rows[0])<=26 or any(not isinstance(r,list) or len(r)!=len(rows[0]) for r in rows):raise ValueError('invalid_spreadsheet')
        names.add(name.casefold());sheet=book.create_sheet(name);sheet.freeze_panes='A2'
        for row in rows:
            converted=[]
            for value in row:
                if isinstance(value,str):
                    value=value.strip()
                    if structured and value.startswith('='):value=checked_formula(value)
                    elif re.fullmatch(r'-?\d+(?:\.\d+)?',value):value=float(value) if '.' in value else int(value)
                    elif value.startswith(('=','+','-','@')):raise ValueError('unsafe_spreadsheet_formula')
                    elif len(value)>2000:raise ValueError('spreadsheet_cell_limit')
                elif value is not None and (not isinstance(value,(int,float,bool)) or isinstance(value,float) and not math.isfinite(value)):raise ValueError('invalid_spreadsheet_cell')
                converted.append(value)
            sheet.append(converted)
        for cell in sheet[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='263244')
        from openpyxl.utils import get_column_letter
        for i in range(1,sheet.max_column+1):sheet.column_dimensions[get_column_letter(i)].width=min(45,max(14,max(len(str(sheet.cell(r,i).value or '')) for r in range(1,sheet.max_row+1))+2))
        sheet.auto_filter.ref=sheet.dimensions
        charts=item.get('charts',[])
        if not isinstance(charts,list) or len(charts)>3:raise ValueError('chart_limit')
        for n,definition in enumerate(charts):
            kind=definition.get('type');column=definition.get('data_column',2);labels=definition.get('label_column',1)
            if kind not in ('bar','line','pie') or type(column) is not int or type(labels) is not int or not 1<=column<=sheet.max_column or not 1<=labels<=sheet.max_column:raise ValueError('invalid_chart')
            chart={'bar':BarChart,'line':LineChart,'pie':PieChart}[kind]();chart.title=str(definition.get('title','Chart'))[:80]
            chart.add_data(Reference(sheet,min_col=column,min_row=1,max_row=sheet.max_row),titles_from_data=True)
            chart.set_categories(Reference(sheet,min_col=labels,min_row=2,max_row=sheet.max_row));sheet.add_chart(chart,'A'+str(sheet.max_row+3+n*16))
    book.calculation=CalcProperties(calcId=191029,fullCalcOnLoad=True,forceFullCalc=True)
    book.save(output)


def render(source, format, title=None):
    title = (title or next((s[2:].strip() for s in source.splitlines() if s.startswith('# ')), 'Hasil pekerjaan'))[:100]
    output = io.BytesIO()
    if format == 'docx':
        from docx import Document
        from docx.shared import Inches, Pt
        doc = Document(); doc.core_properties.title=title; doc.core_properties.author='Kilas Works'
        section=doc.sections[0]; section.top_margin=section.bottom_margin=Inches(.8)
        doc.styles['Normal'].font.name='Calibri'; doc.styles['Normal'].font.size=Pt(11)
        table_rows=[]
        def flush():
            if not table_rows:return
            rows=[row for row in table_rows if not all(re.fullmatch(r'[:\- ]+',c) for c in row)]
            if not rows or any(len(r)!=len(rows[0]) for r in rows):raise ValueError('invalid_office_table')
            table=doc.add_table(rows=0,cols=len(rows[0]));table.style='Light Shading Accent 1'
            for row in rows:
                for cell,text in zip(table.add_row().cells,row):cell.text=text
            table_rows.clear()
        for line in source.splitlines():
            if line.startswith('|'):table_rows.append([v.strip() for v in line.strip('|').split('|')]);continue
            flush()
            if line.startswith('#'):doc.add_heading(line.lstrip('# ').strip(),level=min(3,len(line)-len(line.lstrip('#'))-1))
            elif line.startswith('- '):doc.add_paragraph(line[2:],style='List Bullet')
            elif line.strip():doc.add_paragraph(re.sub(r'\*\*(.*?)\*\*',r'\1',line))
        flush();doc.save(output)
    elif format == 'xlsx':
        spreadsheet(source,output)
    else:
        from pptx import Presentation
        from pptx.util import Inches, Pt
        from pptx.dml.color import RGBColor
        deck=Presentation();deck.slide_width=Inches(13.33);deck.slide_height=Inches(7.5)
        sections=re.split(r'^## ',source,flags=re.M)
        if not 2<=len(sections)<=16:raise ValueError('invalid_presentation_structure')
        for index,section in enumerate(sections):
            lines=[re.sub(r'^[-# ]+','',s).strip() for s in section.splitlines() if s.strip()]
            heading=lines[0][:110];body=lines[1:]
            if len(body)>7 or sum(map(len,body))>1100:raise ValueError('presentation_too_dense')
            slide=deck.slides.add_slide(deck.slide_layouts[6]);slide.background.fill.solid();slide.background.fill.fore_color.rgb=RGBColor(248,249,251)
            box=slide.shapes.add_textbox(Inches(.8),Inches(.7),Inches(11.7),Inches(1.2)).text_frame
            box.text=heading;box.paragraphs[0].font.size=Pt(32);box.paragraphs[0].font.bold=True
            frame=slide.shapes.add_textbox(Inches(.8),Inches(2),Inches(11.7),Inches(4.5)).text_frame;frame.word_wrap=True
            for n,line in enumerate(body):
                p=frame.paragraphs[0] if n==0 else frame.add_paragraph();p.text=line;p.font.size=Pt(21);p.space_after=Pt(16)
            footer=slide.shapes.add_textbox(Inches(.8),Inches(6.9),Inches(10),Inches(.3)).text_frame;footer.text=f'Kilas Works · {index+1}';footer.paragraphs[0].font.size=Pt(10)
        deck.core_properties.title=title;deck.save(output)
    name=re.sub(r'[^a-z0-9-]+','-',title.lower()).strip('-') or 'hasil-pekerjaan'
    return {'filename':name[:70]+'.'+format,'mime_type':work_artifacts.MIMES[format],'content':output.getvalue(),'title':title,'extracted_text':source}

"""Fixed synthetic document acceptance, executed only inside the confined VM."""
import io
import json
from pathlib import Path

PAGES = 50
SLIDES = 15


def documents():
    import random
    from PIL import Image
    from reportlab.pdfgen import canvas
    from reportlab.lib.utils import ImageReader
    from pptx import Presentation
    from pptx.util import Inches
    result = {}
    image = Image.frombytes('RGB', (640, 480), random.Random(19).randbytes(640*480*3))
    png = io.BytesIO()
    image.save(png, format='PNG')
    for name, count in (('small.pdf', 1), ('large.pdf', PAGES)):
        raw = io.BytesIO()
        pdf = canvas.Canvas(raw)
        for index in range(count):
            pdf.drawString(40, 800, f'Synthetic page {index+1}: revenue 410000 cost 226000 profit 184000')
            if count > 1:
                pdf.drawImage(ImageReader(io.BytesIO(png.getvalue())), 40, 300, width=500, height=375)
            pdf.showPage()
        pdf.save()
        result[name] = raw.getvalue()
    for name, count in (('small.pptx', 1), ('large.pptx', SLIDES)):
        deck = Presentation()
        for index in range(count):
            slide = deck.slides.add_slide(deck.slide_layouts[5])
            slide.shapes.title.text = f'Synthetic slide {index+1}: profit 184000'
            if count > 1:
                slide.shapes.add_picture(io.BytesIO(png.getvalue()), Inches(1), Inches(2), width=Inches(6))
        raw = io.BytesIO()
        deck.save(raw)
        result[name] = raw.getvalue()
    return result


def main():
    import boundary_checks as boundary
    stage = 'boundary'
    try:
        boundary.main()
        report = json.loads(Path('output/proof.json').read_text())
        Path('output/proof.json').unlink()
        Path('output/answer.txt').unlink()
        stage = 'documents'
        files = documents()
        for name, raw in files.items():
            Path('output', name).write_bytes(raw)
        print('KILAS_DOCUMENT_OK '+json.dumps({'boundary':report, 'sizes':{n:len(b) for n,b in files.items()},
                                             'pages':PAGES, 'slides':SLIDES}, separators=(',', ':')))
    except BaseException:
        print('KILAS_DOCUMENT_FAILED '+stage)
        raise SystemExit(1)


if __name__ == '__main__':
    main()

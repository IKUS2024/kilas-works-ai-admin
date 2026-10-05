"""Focused Work V2 lifecycle, reminder, Office, push and permission regressions."""
import base64
import io
import json
import os
import unittest
from datetime import datetime, timedelta, timezone
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch
import test_kilas_work_documents as f
from kilas_ai import work_schedule, work_office, work_push, automation_store, work_routes, work_artifacts, agent_store

NOW=datetime(2026,10,2,2,0,tzinfo=timezone.utc)


class WorkV2Tests(unittest.TestCase):
    def client(self,owner=None):
        if owner:return f.WorkTests.client(self,owner)
        if not hasattr(self,'_client'):self._client=f.WorkTests.client(self)
        return self._client
    complete=f.WorkTests.complete

    def setUp(self):
        f.WorkTests.setUp(self)

    def submit(self,text,**extra):
        return self.client().post('/kilas-ai/agent/chat',data={'message':text,'csrf_token':'work-csrf','conversation_id':self.conversation,**extra})

    def tick(self,job):
        f.fixture.db.execute('UPDATE kilas_agent_jobs SET next_wake_at=? WHERE id=?',(f.fixture.store.stamp(),job))
        claims=f.fixture.store.claim_due(1)
        self.assertTrue(claims)
        f.fixture.runner.execute(*claims[0])

    def test_capability_questions_never_create_jobs_or_invoke_connectors(self):
        with patch('kilas_ai.agent_chat.ordinary',return_value=None),patch('kilas_ai.connector_flow.handle',side_effect=AssertionError('connector invoked')):
            for question in ('kamu bisa bikin PDF ga?','Apa itu laporan?','Can you create a PDF?','kamu bisa kirim Gmail?'):
                self.assertEqual(self.submit(question).status_code,303)
        self.assertEqual(f.fixture.store.list_jobs(self.owner),[])

    def test_prestart_clarification_no_job_then_same_conversation_document(self):
        self.submit('buat PDF')
        self.assertEqual(f.fixture.store.list_jobs(self.owner),[])
        self.submit('company profile Kilas Works')
        jobs=f.fixture.store.list_jobs(self.owner)
        self.assertEqual(len(jobs),1)
        self.assertIn('company profile',jobs[0]['instruction'])
        self.assertEqual(jobs[0]['origin_conversation_id'],self.conversation)

    def test_started_clarification_waits_without_artifact_and_reply_resumes_same_job(self):
        self.submit(f.REQUEST)
        job=f.fixture.store.list_jobs(self.owner)[0]['id']
        with patch.object(f.content_worker,'text',return_value=('Apa yang perlu dicantumkan dalam lingkup pekerjaan?',f.model_policy.LUNA,{})):
            self.tick(job);self.tick(job)
        stored=f.fixture.store.get(self.owner,job)
        self.assertEqual(stored['status'],'WAITING')
        self.assertEqual(stored['last_error'],'waiting_input')
        self.assertIsNone(stored['next_wake_at'])
        self.assertFalse(work_artifacts.listing(self.owner,job_id=job))
        self.assertIn('Menunggu jawabanmu',self.client().get('/kilas-ai/agent').text)
        self.submit('Lingkupnya mencakup layanan restoran')
        self.assertEqual(len(f.fixture.store.list_jobs(self.owner)),1)
        with patch.object(f.content_worker,'text',return_value=(f.SOURCE,f.model_policy.LUNA,{})):
            self.tick(job);self.tick(job)
        self.assertEqual(f.fixture.store.get(self.owner,job)['status'],'COMPLETED')

    def test_document_is_persisted_before_owner_scoped_bounded_start(self):
        self.submit(f.REQUEST)
        job=f.fixture.store.list_jobs(self.owner)[0]['id']
        path=f'/kilas-ai/work/jobs/{job}/start'
        self.assertFalse(self.client(self.other).post(path,json={},headers={'X-CSRF-Token':'work-csrf'}).json['started'])
        with patch.object(f.content_worker,'text',return_value=(f.SOURCE,f.model_policy.LUNA,{})):
            self.assertTrue(self.client().post(path,json={},headers={'X-CSRF-Token':'work-csrf'}).json['started'])
            self.assertTrue(self.client().post(path,json={},headers={'X-CSRF-Token':'work-csrf'}).json['started'])
        self.assertEqual(f.fixture.store.get(self.owner,job)['status'],'COMPLETED')
        self.assertFalse(self.client().post(path,json={},headers={'X-CSRF-Token':'work-csrf'}).json['started'])
        self.assertEqual(self.client().post(path).status_code,400)
        self.assertEqual(len(work_artifacts.listing(self.owner,job_id=job)),1)

    def test_active_count_not_unread_events_and_new_work_keeps_background_job(self):
        active=f.fixture.store.create(self.owner,'Riset terbaru',conversation_id=self.conversation)
        completed=f.fixture.store.create(self.owner,'Pekerjaan lama',conversation_id=self.conversation)
        f.fixture.db.execute("UPDATE kilas_agent_jobs SET status='COMPLETED' WHERE id=?",(completed,))
        with f.fixture.store.transaction() as conn:
            for _ in range(6):f.fixture.store.event(conn,completed,'FAILED','Internal diagnostic')
            f.fixture.store.event(conn,active,'CONDITION_MET','Kondisi terpantau terpenuhi.')
        html=self.client().get('/kilas-ai/agent').text
        self.assertNotIn('data-active-count',html)
        self.assertNotIn('Pekerjaan aktif',html)
        self.assertNotIn('view=notifications',html)
        self.assertIn(f'data-job-id="{active}"',html)
        self.assertNotIn('>Connections</a>',html)
        self.assertNotIn('Advanced settings',html)
        self.assertNotIn('autonomous_mode',html)
        self.assertIn('Kondisi terpantau terpenuhi.',self.client().get('/kilas-ai/agent?view=notifications').text)
        self.client().post('/kilas-ai/agent/conversations',data={'csrf_token':'work-csrf'})
        self.assertEqual(f.fixture.store.get(self.owner,active)['status'],'PLANNING')

    def test_completed_result_stays_in_original_chat_turn_after_open_and_new_message(self):
        agent_store.append(self.owner,'user',f.REQUEST,self.conversation)
        job,file=self.complete(f.REQUEST,f.SOURCE)
        self.client().get(f'/kilas-ai/agent/jobs/{job}')
        self.client().get(f"/kilas-ai/agent/jobs/{job}/artifacts/{file['id']}?download=1")
        with patch('kilas_ai.agent_chat.ordinary',return_value=None):self.submit('Pesan baru sesudah hasil lama')
        for _ in range(2):
            chat=self.client().get('/kilas-ai/agent').text
            self.assertEqual(chat.count(f'data-job-id="{job}"'),1)
            self.assertIn(file['name'],chat)
            self.assertLess(chat.index(f'data-job-id="{job}"'),chat.index('>Pesan baru sesudah hasil lama<'))
            self.assertNotIn('Selesai terbaru',chat)
        self.assertIn(f'data-job-id="{job}"',self.client().get('/kilas-ai/agent?view=history').text)
        self.assertEqual(f.fixture.store.get(self.owner,job)['status'],'COMPLETED')

    def test_old_results_beyond_eight_survive_acknowledgment_and_message_pagination(self):
        original=[]
        for i in range(12):
            agent_store.append(self.owner,'user',f'Buat PDF arsip {i}',self.conversation)
            job=f.fixture.store.create(self.owner,f'Buat PDF arsip {i}',conversation_id=self.conversation)
            f.fixture.db.execute("UPDATE kilas_agent_jobs SET status='COMPLETED' WHERE id=?",(job,))
            original.append(job)
        # Legacy data has no saved anchor, and its created_at may be a native PG datetime.
        f.fixture.db.execute("UPDATE kilas_agent_jobs SET checkpoint_json='{}' WHERE id=?",(original[0],))
        chat=self.client().get('/kilas-ai/agent').text
        for job in original:self.assertEqual(chat.count(f'data-job-id="{job}"'),1)
        for i in range(65):agent_store.append(self.owner,'user',f'Pesan berikutnya {i}',self.conversation)
        newest=agent_store.messages(self.owner,conversation_id=self.conversation)
        chat=self.client().get('/kilas-ai/agent').text
        for job in original:self.assertNotIn(f'data-job-id="{job}"',chat)
        earlier=self.client().get(f"/kilas-ai/agent?before={newest[0]['id']}").text
        for job in original:self.assertEqual(earlier.count(f'data-job-id="{job}"'),1)
        self.assertNotIn(f'data-job-id="{original[0]}"',self.client(self.other).get('/kilas-ai/agent').text)

    def test_submissions_anchor_to_their_own_message_even_when_another_arrives(self):
        first=agent_store.append(self.owner,'user','Buat PDF pertama',self.conversation)
        second=agent_store.append(self.owner,'user','Buat PDF kedua',self.conversation)
        job=f.fixture.store.create(self.owner,'Buat PDF pertama',conversation_id=self.conversation,origin_message_id=first)
        self.assertEqual([j['id'] for j in f.fixture.store.conversation_jobs(self.owner,self.conversation,[first])],[job])
        self.assertEqual(f.fixture.store.conversation_jobs(self.owner,self.conversation,[second]),[])
        foreign=agent_store.new_conversation(self.other)
        foreign_message=agent_store.append(self.other,'user','Buat PDF luar',foreign)
        with self.assertRaisesRegex(ValueError,'message_not_owned'):
            f.fixture.store.create(self.owner,'Buat PDF pertama',conversation_id=self.conversation,origin_message_id=foreign_message)

    def test_informal_pdf_request_creates_real_document_instead_of_chat_prose(self):
        request='bikinin gw pdf dah bro untuk company profile random aja untuk kilasworks yang bergerak di bidang sosial media'
        with patch('kilas_ai.agent_chat.ordinary',side_effect=AssertionError('PDF must use real document pipeline')):
            self.submit(request)
        job=f.fixture.store.list_jobs(self.owner)[0]['id']
        with patch.object(f.content_worker,'text',return_value=('# Company Profile Kilas Works\n\n## Tentang Kami\nKilas Works adalah brand sintetis layanan sosial media.\n\n## Layanan\nPerencanaan konten dan pengelolaan media sosial.',f.model_policy.LUNA,{})):
            self.tick(job);self.tick(job)
        artifact=work_artifacts.listing(self.owner,job_id=job)[0]
        response=self.client().get(f"/kilas-ai/agent/jobs/{job}/artifacts/{artifact['id']}")
        self.assertTrue(response.data.startswith(b'%PDF'))
        self.assertEqual(f.fixture.store.get(self.owner,job)['status'],'COMPLETED')

    def test_raw_execution_prompts_not_visible_and_progress_persists(self):
        job=f.fixture.store.create(self.owner,'Buat proposal',conversation_id=self.conversation)
        claim=f.fixture.store.claim_due(1)[0]
        plan=f.fixture.planner.validate(f.fixture.proposal(),'ONE_SHOT')
        plan['steps'][0]['instruction']='SECRET_RAW_WORKER_PROMPT'
        f.fixture.store.install_plan(f.fixture.store.get(self.owner,job),claim[1],plan)
        with f.fixture.store.transaction() as conn:f.fixture.store.event(conn,job,'WRITING','SECRET_PROVIDER_WORDING',unread=False)
        f.fixture.store.release(*claim)
        for path in ('/kilas-ai/agent',f'/kilas-ai/agent/jobs/{job}'):
            html=self.client().get(path).text
            self.assertNotIn('SECRET_RAW_WORKER_PROMPT',html)
            self.assertNotIn('SECRET_PROVIDER_WORDING',html)
            self.assertNotIn('Instruksi pelaksanaan',html)
        self.assertIn('Menulis isi',self.client().get('/kilas-ai/agent').text)

    def test_timezone_browser_updates_existing_setting_and_authoritative_clock(self):
        with patch('kilas_ai.agent_chat.ordinary',return_value=None):self.submit('Halo',browser_timezone='Asia/Bangkok')
        self.assertEqual(automation_store.setting(self.owner),'Asia/Bangkok')
        path='/kilas-ai/work/preferences'
        response=self.client().post(path,json={'timezone':'Asia/Jayapura','manual':True},headers={'X-CSRF-Token':'work-csrf'})
        self.assertEqual(response.status_code,200)
        self.client().post(path,json={'timezone':'Europe/London'},headers={'X-CSRF-Token':'work-csrf'})
        self.assertEqual(automation_store.setting(self.owner),'Europe/London')
        with patch.object(f.fixture.store,'now',return_value=NOW):self.submit('jam berapa sekarang')
        self.assertIn('03.00',agent_store.messages(self.owner,conversation_id=self.conversation)[-1]['content'])
        invalid=self.client().post(path,json={'timezone':'Invalid/Place'},headers={'X-CSRF-Token':'work-csrf'})
        self.assertEqual(invalid.status_code,400)
        self.assertEqual(automation_store.setting(self.owner),'Europe/London')
        self.assertNotIn('id="work-timezone"',self.client().get('/kilas-ai/agent?view=settings').text)

    def test_location_required_only_for_relevant_request_and_bounded_fresh_payload(self):
        self.submit('Cari restoran dekat sini')
        self.assertFalse(f.fixture.store.list_jobs(self.owner))
        self.assertIn('Izinkan lokasi',self.client().get('/kilas-ai/agent').text)
        with patch.object(f.fixture.store,'now',return_value=NOW):
            valid={'permission_granted':True,'latitude':-6.2,'longitude':106.8,'accuracy':40,'timestamp':NOW.timestamp()*1000}
            self.assertEqual(work_routes.location_payload(json.dumps(valid))['accuracy'],40)
            for change in ({'latitude':100},{'timestamp':0},{'permission_granted':False},{'longitude':'NaN'}):
                with self.assertRaises(ValueError):work_routes.location_payload(json.dumps({**valid,**change}))
        self.submit('Jakarta Selatan')
        self.assertIn('Jakarta Selatan',agent_store.messages(self.owner,conversation_id=self.conversation)[-1]['content'] if not f.fixture.store.list_jobs(self.owner) else f.fixture.store.list_jobs(self.owner)[0]['instruction'])

    def test_reminder_subject_clarification_local_schedule_no_gps_or_model(self):
        automation_store.set_timezone(self.owner,'Asia/Bangkok')
        with patch.object(f.fixture.store,'now',return_value=NOW),patch.object(f.fixture.planner.requests,'post',side_effect=AssertionError('No model')):
            self.submit('ingatkan aku besok jam 8')
            self.assertFalse(f.fixture.store.list_jobs(self.owner))
            self.submit('bayar tagihan')
        jobs=f.fixture.store.list_jobs(self.owner)
        self.assertEqual(len(jobs),1)
        self.assertEqual(jobs[0]['next_wake_at'],'2026-10-03T01:00:00+00:00')
        self.assertEqual(json.loads(jobs[0]['checkpoint_json'])['reminder']['subject'],'bayar tagihan')
        self.assertNotIn('data-work-location',self.client().get('/kilas-ai/agent').text)
        self.assertFalse(work_push.configured())
        self.assertIn('Notifikasi perangkat belum aktif',agent_store.messages(self.owner,conversation_id=self.conversation)[-1]['content'])

    def test_frozen_time_relative_weekly_monthly_dates_and_past(self):
        self.assertEqual(work_schedule.subject('ingatkan setiap hari jam 8 minum air'),'minum air')
        self.assertEqual(work_schedule.subject('ingatkan tanggal 10 Oktober jam 9 bayar tagihan'),'bayar tagihan')
        samples=(('ingatkan minum air 20 menit lagi','Asia/Jakarta',NOW+timedelta(minutes=20)),('ingatkan minum air 3 jam lagi','Asia/Bangkok',NOW+timedelta(hours=3)),('ingatkan besok jam 8 bayar tagihan','Asia/Bangkok',datetime(2026,10,3,1,tzinfo=timezone.utc)))
        for text,zone,want in samples:self.assertEqual(work_schedule.parse(text,zone,NOW)['next_run_at'],want)
        self.assertEqual(work_schedule.parse('ingatkan tiap Senin jam 8 cek jadwal','Asia/Jakarta',NOW)['schedule']['kind'],'weekly')
        self.assertEqual(work_schedule.parse('ingatkan tiap tanggal 10 jam 9 bayar','Asia/Jakarta',NOW)['schedule']['kind'],'monthly')
        self.assertEqual(work_schedule.parse('ingatkan tanggal 10 jam 9 bayar','Asia/Jakarta',NOW)['next_run_at'].day,10)
        self.assertEqual(work_schedule.parse('ingatkan bulan depan jam 9 bayar','Asia/Jakarta',NOW)['next_run_at'].month,11)
        with self.assertRaises(ValueError):work_schedule.parse('ingatkan hari ini jam 8 minum air','Asia/Jakarta',NOW)
        with self.assertRaises(ValueError):work_schedule.parse('ingatkan tanggal 1/10/2026 jam 8 minum air','Asia/Jakarta',NOW)

    def test_delayed_reminder_same_conversation_once_unread_with_delivery_times(self):
        with patch.object(f.fixture.store,'now',return_value=NOW):self.submit('ingatkan aku 20 menit lagi minum air')
        job=f.fixture.store.list_jobs(self.owner)[0]['id']
        later=NOW+timedelta(hours=2)
        with patch.object(f.fixture.store,'now',return_value=later),patch.object(f.fixture.planner.requests,'post',side_effect=AssertionError('No model')):
            for _ in range(2):f.fixture.runner.execute(*f.fixture.store.claim_due(1)[0])
            self.assertEqual(f.fixture.store.claim_due(1),[])
        messages=agent_store.messages(self.owner,conversation_id=self.conversation)
        self.assertEqual(sum(m['content']=='Pengingat: minum air' for m in messages),1)
        event=f.fixture.db.query_one("SELECT * FROM kilas_agent_events WHERE job_id=? AND kind='REMINDER'",(job,))
        self.assertEqual(event['unread'],1)
        data=json.loads(f.fixture.store.get(self.owner,job)['checkpoint_json'])['reminder']
        self.assertEqual(data['scheduled_for'],f.fixture.store.stamp(NOW+timedelta(minutes=20)))
        self.assertEqual(data['actual_delivery_time'],f.fixture.store.stamp(later))

    def test_recurring_reminder_more_than_step_cap_without_duplicate_or_model(self):
        with patch.object(f.fixture.store,'now',return_value=NOW):self.submit('ingatkan tiap Senin jam 8 cek jadwal')
        job=f.fixture.store.list_jobs(self.owner)[0]['id']
        with patch.object(f.fixture.planner.requests,'post',side_effect=AssertionError('No model')):
            for cycle in range(26):
                wake=f.fixture.usage._as_utc(f.fixture.store.get(self.owner,job)['next_wake_at'])
                with patch.object(f.fixture.store,'now',return_value=wake+timedelta(seconds=1)):
                    if cycle==0:f.fixture.runner.execute(*f.fixture.store.claim_due(1)[0])
                    f.fixture.runner.execute(*f.fixture.store.claim_due(1)[0])
        self.assertEqual(len(f.fixture.store.steps(job)),1)
        self.assertEqual(f.fixture.store.get(self.owner,job)['cycle'],26)
        self.assertEqual(f.fixture.db.query_one("SELECT COUNT(*) AS n FROM kilas_agent_events WHERE job_id=? AND kind='REMINDER'",(job,))['n'],26)

    def test_natural_reminder_edit_retains_job_then_cancel(self):
        with patch.object(f.fixture.store,'now',return_value=NOW):
            self.submit('ingatkan besok jam 8 bayar tagihan')
            job=f.fixture.store.list_jobs(self.owner)[0]['id']
            self.submit('ubah reminder tadi jadi jam 9')
            self.assertEqual(f.fixture.store.get(self.owner,job)['next_wake_at'],'2026-10-03T02:00:00+00:00')
            self.assertEqual(json.loads(f.fixture.store.get(self.owner,job)['checkpoint_json'])['reminder']['subject'],'bayar tagihan')
            self.submit('ulang tiap Senin')
            self.assertEqual(f.fixture.store.get(self.owner,job)['mode'],'RECURRING')
            self.submit('batalin reminder besok')
            self.assertEqual(f.fixture.store.get(self.owner,job)['status'],'STOPPED')
        self.assertEqual(len(f.fixture.store.list_jobs(self.owner)),1)

    def test_real_docx_xlsx_pptx_and_formula_safety(self):
        sources={'docx':f.SOURCE,'xlsx':'Item,Jumlah\nLunch box,30\nAir,2\n','pptx':'# Kilas Works\nSolusi untuk UMKM\n## Layanan\n- Jawaban customer\n- Pekerjaan operasional\n## Langkah Berikutnya\n- Konfirmasi lingkup\n'}
        for format,source in sources.items():
            job,file=self.complete(f.REQUEST+' '+format if format=='docx' else 'Buat dokumen '+format,source)
            response=self.client().get(f"/kilas-ai/agent/jobs/{job}/artifacts/{file['id']}")
            self.assertTrue(response.data.startswith(b'PK'))
            work_artifacts.validate({'filename':file['name'],'mime_type':file['media_type'],'content':response.data})
            if format=='docx':
                from docx import Document
                self.assertIn('Proposal Kerja Sama',[p.text for p in Document(io.BytesIO(response.data)).paragraphs])
            elif format=='xlsx':
                from openpyxl import load_workbook
                book=load_workbook(io.BytesIO(response.data));self.assertEqual(book.active['B2'].value,30);book.close()
            else:
                from pptx import Presentation
                self.assertEqual(len(Presentation(io.BytesIO(response.data)).slides),3)
        with self.assertRaises(ValueError):work_office.render('Name,Value\nBad,=WEBSERVICE("http://evil")','xlsx')

    def test_office_source_from_composer_and_image_edit_owned_source(self):
        sheet=work_office.render('Jumlah,Harga\n0,0','xlsx')
        self.submit('Buat proposal PDF dari dokumen ini',source_files=(io.BytesIO(sheet['content']),'source.xlsx',sheet['mime_type']))
        sheet_job=f.fixture.store.list_jobs(self.owner)[0]
        self.assertIn('A2=0 | B2=0',json.loads(sheet_job['checkpoint_json'])['source_materials'][0]['text'])
        f.fixture.store.control(self.owner,sheet_job['id'],'stop')
        file=work_office.render('# Source\n\n## Facts\nVerified facts about Kilas Works.','docx')
        self.submit('Buat proposal PDF dari dokumen ini',source_files=(io.BytesIO(file['content']),'source.docx',file['mime_type']))
        job=f.fixture.store.list_jobs(self.owner)[0]
        self.assertIn('Verified facts',json.loads(job['checkpoint_json'])['source_materials'][0]['text'])
        f.fixture.store.control(self.owner,job['id'],'stop')
        from PIL import Image
        image=io.BytesIO();Image.new('RGB',(24,24),'orange').save(image,'PNG')
        self.submit('Ubah gambar ini jadi lebih terang',source_files=(io.BytesIO(image.getvalue()),'photo.png','image/png'))
        job=f.fixture.store.list_jobs(self.owner)[0]
        source=json.loads(job['checkpoint_json'])['image_source_id']
        with self.assertRaises(ValueError):work_artifacts.image_source(self.other,source,self.conversation)
        self.assertEqual(work_artifacts.listing(self.owner,job_id=job['id']),[])
        with patch.dict(os.environ,{'KILAS_AI_OPENAI_IMAGE_MODEL':'gpt-image-2'}),patch('kilas_ai.tools.image',return_value={'raw':image.getvalue(),'mime':'image/png','model':'gpt-image-2','usage':{}}) as provider:
            self.tick(job['id']);self.tick(job['id'])
        self.assertEqual(provider.call_args.kwargs['source']['content'],image.getvalue())
        self.assertEqual(f.fixture.store.get(self.owner,job['id'])['status'],'COMPLETED')

    def subscription(self):
        from cryptography.hazmat.primitives.asymmetric import ec
        from cryptography.hazmat.primitives import serialization
        point=ec.generate_private_key(ec.SECP256R1()).public_key().public_bytes(serialization.Encoding.X962,serialization.PublicFormat.UncompressedPoint)
        encode=lambda raw:base64.urlsafe_b64encode(raw).decode().rstrip('=')
        return {'endpoint':'https://fcm.googleapis.com/fcm/send/test-'+str(self.owner),'keys':{'p256dh':encode(point),'auth':encode(b'1234567890123456')}}

    def test_push_owner_isolation_validation_csrf_https_and_unconfigured_honesty(self):
        data=self.subscription();sub=work_push.register(self.owner,data)
        self.assertEqual(work_push.register(self.owner,data),sub)
        with self.assertRaises(ValueError):work_push.register(self.other,data)
        for endpoint in ('http://localhost/x','https://127.0.0.1/x','https://evil.test/x','https://fcm.googleapis.com:8443/x'):
            with self.assertRaises(ValueError):work_push.register(self.owner,{**data,'endpoint':endpoint})
        with self.assertRaises(ValueError):work_push.register(self.owner,[])
        self.assertEqual(self.client().post('/kilas-ai/work/push',json=data).status_code,400)
        self.assertEqual(self.client().post('/kilas-ai/work/push',json=data,headers={'X-CSRF-Token':'work-csrf'}).status_code,400)
        self.client(self.other).delete(f'/kilas-ai/work/push/{sub}',json={},headers={'X-CSRF-Token':'work-csrf'})
        self.assertIsNotNone(f.fixture.db.query_one('SELECT id FROM kilas_work_push_subscriptions WHERE id=?',(sub,)))

    def test_push_at_most_once_delivery_and_invalid_subscription_disabled(self):
        data=self.subscription();sub=work_push.register(self.owner,data)
        config={'KILAS_WEB_PUSH_PUBLIC_KEY':'qa-public','KILAS_WEB_PUSH_PRIVATE_KEY':'qa-private','KILAS_WEB_PUSH_SUBJECT':'mailto:qa@example.test'}
        with patch.dict(os.environ,config),patch.object(f.fixture.store,'now',return_value=NOW):self.submit('ingatkan 20 menit lagi minum air')
        job=f.fixture.store.list_jobs(self.owner)[0]['id']
        with patch.dict(os.environ,config):
            self.tick(job);self.tick(job)
            with patch('pywebpush.webpush') as transport:
                self.assertEqual(work_push.deliver(),1)
                self.assertEqual(work_push.deliver(),0)
                self.assertEqual(transport.call_count,1)
                payload=json.loads(transport.call_args.args[1])
                self.assertIn(str(self.conversation),payload['url'])
            f.fixture.db.execute("UPDATE kilas_work_push_deliveries SET state='PENDING' WHERE subscription_id=?",(sub,))
            from pywebpush import WebPushException
            with patch('pywebpush.webpush',side_effect=WebPushException('Expired',response=SimpleNamespace(status_code=410))):work_push.deliver()
        self.assertEqual(f.fixture.db.query_one('SELECT enabled FROM kilas_work_push_subscriptions WHERE id=?',(sub,))['enabled'],0)

    def test_sqlite_push_migration_idempotency_and_no_unrelated_changes(self):
        path=Path(__file__).parents[1]/'migrations/0081_kilas_work_push_sqlite.sql'
        with f.fixture.store.transaction() as conn:
            conn.executescript(path.read_text());conn.executescript(path.read_text())
        self.assertIsNotNone(f.fixture.db.query_one("SELECT name FROM sqlite_master WHERE name='kilas_work_push_subscriptions'"))


if __name__=='__main__':unittest.main()

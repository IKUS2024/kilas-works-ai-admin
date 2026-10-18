"""Large local fixtures and real SDK stream objects; never create a remote VM."""
import ast,hashlib,io,json,os,random,tempfile,time,unittest,zipfile
from pathlib import Path
from unittest.mock import Mock,patch
import httpx
from e2b.sandbox.filesystem.filesystem import FileStreamReader
from kilas_ai import workspace_large_outputs as large,workspace_large_guest as guest,workspace_large_execution as remote
from kilas_ai import workspace_output_formats as formats,workspace_execution_worker as worker,workspace_projects as projects
from kilas_ai import workspace_execution_proof as proof,workspace_execution_rollout as rollout,work_artifacts,autonomous_runner as runner
import test_kilas_workspace_execution as f
import db


def manifest(files):
    return {'protocol':large.PROFILE,'files':[{'name':name,'slot':f'{i:03}','size':len(raw),'sha256':hashlib.sha256(raw).hexdigest()} for i,(name,raw) in enumerate(files.items())]}


class Body(httpx.SyncByteStream):
    def __init__(self,raw,fail=False):self.raw=raw;self.closed=False;self.fail=fail
    def __iter__(self):
        for pos in range(0,len(self.raw),65536):
            yield self.raw[pos:pos+65536]
            if self.fail:raise httpx.ReadError('synthetic interrupted stream')
    def close(self):self.closed=True


def sandbox(files,fail=False):
    vm=Mock();bodies=[]
    def read(path,**kw):
        raw=list(files.values())[int(path.rsplit('/',1)[1])];body=Body(raw,fail);bodies.append(body)
        return FileStreamReader(httpx.Response(200,headers={'content-length':str(len(raw))},stream=body))
    vm.files.read.side_effect=read
    return vm,bodies


def fixtures():
    rng=random.Random(18);text=''.join(rng.choices('abcdefghijklmnopqrstuvwxyz0123456789 ',k=200000))
    from docx import Document
    doc=Document();doc.add_paragraph(text);out=io.BytesIO();doc.save(out);files={'large.docx':out.getvalue()}
    from openpyxl import Workbook
    book=Workbook();sheet=book.active
    for i in range(10000):sheet.append([i,text[i:i+80],f'=A{i+1}*2'])
    out=io.BytesIO();book.save(out);files['large.xlsx']=out.getvalue()
    from PIL import Image
    image=Image.frombytes('RGB',(640,480),rng.randbytes(640*480*3));png=io.BytesIO();image.save(png,format='PNG')
    from pptx import Presentation
    deck=Presentation();slide=deck.slides.add_slide(deck.slide_layouts[6]);slide.shapes.add_picture(io.BytesIO(png.getvalue()),0,0);out=io.BytesIO();deck.save(out);files['large.pptx']=out.getvalue()
    from reportlab.pdfgen import canvas
    from reportlab.lib.utils import ImageReader
    out=io.BytesIO();pdf=canvas.Canvas(out);pdf.drawImage(ImageReader(io.BytesIO(png.getvalue())),20,300,width=500,height=400);pdf.save();files['large.pdf']=out.getvalue()
    return files


class LargeTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):cls.files=fixtures()
    def test_real_large_office_pdf_stream_hash_and_format(self):
        self.assertTrue(all(65536<len(raw)<large.MAX_FILE for raw in self.files.values()))
        vm,bodies=sandbox(self.files);result=large.download(vm,manifest(self.files),time.monotonic()+30)
        self.assertEqual({f['filename']:f['content'] for f in result},self.files);self.assertTrue(all(b.closed for b in bodies))
        for call in vm.files.read.call_args_list:self.assertEqual(call.kwargs['format'],'stream');self.assertFalse(call.kwargs['gzip'])
    def test_exact_file_and_job_limits_and_oversize_reject_before_read(self):
        files={'one.txt':b'a'*large.MAX_FILE,'two.txt':b'b'*large.MAX_FILE};vm,bodies=sandbox(files)
        result=large.download(vm,manifest(files),time.monotonic()+30);self.assertEqual(sum(len(f['content']) for f in result),large.MAX_JOB)
        for invalid in ({'too.txt':b'a'*(large.MAX_FILE+1)},{**files,'extra.txt':b'x'}):
            vm,_=sandbox(invalid)
            with self.assertRaises(ValueError):large.download(vm,manifest(invalid),time.monotonic()+30)
            vm.files.read.assert_not_called()
    def test_stream_interrupt_timeout_cancel_tamper_close_and_remove_temp(self):
        files={'large.pdf':self.files['large.pdf']}
        with tempfile.TemporaryDirectory() as folder,patch('tempfile.tempdir',folder):
            vm,bodies=sandbox(files,True)
            with self.assertRaises(httpx.ReadError):large.download(vm,manifest(files),time.monotonic()+10)
            self.assertTrue(bodies[0].closed);self.assertEqual(list(Path(folder).iterdir()),[])
            for invalid in ('hash','cancel','deadline','size'):
                vm,bodies=sandbox(files);value=manifest(files)
                if invalid=='hash':value['files'][0]['sha256']='0'*64
                if invalid=='size':value['files'][0]['size']-=1
                with self.assertRaises(ValueError):large.download(vm,value,time.monotonic()+(-1 if invalid=='deadline' else 10),active=lambda:invalid!='cancel')
                self.assertTrue(all(b.closed for b in bodies));self.assertEqual(list(Path(folder).iterdir()),[])
    def test_manifest_paths_collisions_schema_and_unknown_stream_fail_closed(self):
        files={'ok.txt':b'ok'};value=manifest(files)
        for change in ({'name':'../escape.pdf'},{'name':'program.exe'},{'slot':'../root'},{'size':True},{'sha256':'bad'},{'extra':1}):
            invalid=json.loads(json.dumps(value));invalid['files'][0].update(change)
            with self.assertRaises(ValueError):large.manifest(invalid)
        invalid=manifest({'a.txt':b'a','A.txt':b'b'})
        with self.assertRaises(ValueError):large.manifest(invalid)
        from contextlib import nullcontext
        vm=Mock();vm.files.read.return_value=nullcontext(object())
        with self.assertRaisesRegex(ValueError,'stream_contract'):large.download(vm,value,time.monotonic()+10)
    def test_trusted_staging_and_symlink_overflow_permissions(self):
        with tempfile.TemporaryDirectory() as folder,patch.object(guest,'UID',os.getuid(),create=True):
            root=Path(folder);output=root/'output';output.mkdir();raw=self.files['large.pdf'];(output/'large.pdf').write_bytes(raw)
            result=guest.stage_outputs(output,root/'private');self.assertEqual(result,manifest({'large.pdf':raw}));self.assertEqual((root/'private'/'000').read_bytes(),raw)
            self.assertEqual((root/'private').stat().st_mode&0o777,0o700);self.assertEqual((root/'private'/'000').stat().st_mode&0o777,0o600)
            (output/'large.pdf').unlink();(output/'bad.pdf').symlink_to(root/'private'/'000')
            with self.assertRaises(OSError):guest.stage_outputs(output,root/'other')
            (output/'bad.pdf').unlink()
            with (output/'big.txt').open('wb') as file:file.truncate(large.MAX_FILE+1)
            with self.assertRaisesRegex(ValueError,'unsafe_output'):guest.stage_outputs(output,root/'third')
    def test_profile_gate_before_model_reservation_sdk_and_original_proof_unchanged(self):
        with patch('e2b.Sandbox.create') as create,patch.object(remote,'reserve') as reserve,patch.object(projects.native,'turn') as model:
            with self.assertRaisesRegex(ValueError,'large_output_acceptance_pending'):remote.execute({'user_id':9},{},{'main.py':'print(42)'})
            self.assertEqual(worker.run({'user_id':9},{},{'request':'Run Python large PDF project'}) .status,'WAITING_CAPABILITY')
            self.assertEqual(projects.run({'user_id':9},{},{'request':'Create large PDF project'}).status,'WAITING_CAPABILITY')
            create.assert_not_called();reserve.assert_not_called();model.assert_not_called()
        self.assertFalse(large.capability(9)['available']);self.assertEqual(proof.fingerprint(),'84f9744af2820215762ee4d60bcc9db560ef5a2ed1176cc8448fe2999e3dbd7b')
        self.assertEqual(hashlib.sha256(Path(rollout.__file__).read_bytes()).hexdigest(),'1403ac10aba25b6fa29916991add1f32e143d6a8cb59b25153e0ef70e4336d83')
        source=guest.runner_source();ast.parse(source)
        self.assertIn("('RLIMIT_FSIZE',8388608)",source);self.assertIn("('RLIMIT_AS',256*1024*1024)",source)
        self.assertLess(source.index('kill_children();process.wait'),source.index("report['outputs']=stage_outputs()"))
    def test_large_zip_expansion_bound_without_weakening_small_archive(self):
        out=io.BytesIO();raw=random.Random(2).randbytes(3*1024*1024).hex().encode()
        with zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED) as archive:archive.writestr('data.txt',raw)
        formats.validate('project.zip','application/zip',out.getvalue(),large=True)
        with self.assertRaisesRegex(ValueError,'archive_size'):formats.validate('project.zip','application/zip',out.getvalue())
    def test_postgres_storage_lock_precedes_capacity_reads(self):
        conn=object();calls=[]
        def query(connection,sql,params,**kw):
            self.assertIs(connection,conn);calls.append(sql)
            return (0,) if sql.startswith('SELECT COALESCE') else None
        with patch.object(db,'BACKEND','postgres'),patch.object(large.usage,'_query',side_effect=query),patch.object(large.usage,'_rows',return_value=[]):
            large.storage_check(conn,large.STORAGE_RESERVATION,require_idle=True)
        self.assertIn('pg_advisory_xact_lock',calls[0]);self.assertIn('SUM(byte_size)',calls[1])


class StorageTests(unittest.TestCase):
    def setUp(self):
        import uuid
        identity=self.id();self.id=lambda:identity+'-'+uuid.uuid4().hex[:8]
        f.ExecutionTests.setUp(self)
    claimed=f.ExecutionTests.claimed
    client=f.ExecutionTests.client
    def enabled(self,vm):
        value=f.ExecutionTests.enabled(self,vm)
        self.addCleanup(patch.stopall)
        patch.object(large,'host_memory',return_value={'limit':536870912,'current':104857600,'peak':209715200}).start()
        return value
    def step(self,job):
        db.execute("INSERT INTO kilas_agent_steps(job_id,sequence,worker,action,instruction,input_json,idempotency_key,status) VALUES (?,1,'EXECUTION','python','Synthetic','{}',?,'RUNNING')",(job['id'],'large-'+str(job['id'])))
        return dict(db.query_one('SELECT * FROM kilas_agent_steps WHERE job_id=?',(job['id'],)))
    def test_durable_storage_exposure_and_atomic_cap_rollback(self):
        from kilas_ai import autonomous_store as jobs,usage
        job,_=self.claimed();step=self.step(job)
        db.execute("INSERT INTO kilas_workspace_executions(user_id,job_id,operation_key,state,reserved_micro_usd,created_at,expires_at,outcome_json,cleanup_confirmed) VALUES (?,?,?,'SUCCEEDED',8880,?,?,?,1)",(job['user_id'],job['id'],'synthetic-large',jobs.stamp(),jobs.stamp(),json.dumps({'large_storage_reserved':large.STORAGE_RESERVATION})))
        run=db.query_one('SELECT id FROM kilas_workspace_executions')['id']
        db.execute("INSERT INTO kilas_workspace_executions(user_id,job_id,operation_key,state,reserved_micro_usd,created_at,expires_at) VALUES (?,?,?,'CREATING',8880,?,?)",(job['user_id'],job['id'],'synthetic-second',jobs.stamp(),jobs.stamp()))
        second=db.query_one('SELECT id FROM kilas_workspace_executions WHERE operation_key=?',('synthetic-second',))['id']
        with self.assertRaisesRegex(ValueError,'storage_busy'):large.reserve_storage(second)
        with jobs.transaction() as conn:
            used=usage._query(conn,'SELECT COALESCE(SUM(byte_size),0) FROM kilas_agent_artifact_files',(),one=True)[0]
            with patch.object(large,'STORAGE_CEILING',used+large.STORAGE_RESERVATION):
                with self.assertRaisesRegex(ValueError,'storage_full'):large.storage_check(conn,1)
                large.storage_check(conn,1,exclude_run=run)
        raw=LargeTests.files['large.pdf'];file={'filename':'large.pdf','mime_type':'application/pdf','content':raw,'validation':{'transport':large.PROFILE,'large_run_id':run}}
        with self.assertRaisesRegex(ValueError,'storage_full'):
            with jobs.transaction() as conn,patch.object(large,'STORAGE_CEILING',used+len(raw)):
                work_artifacts.persist(conn,job,step,file);work_artifacts.persist(conn,job,step,{**file,'filename':'other.pdf'})
        self.assertEqual(work_artifacts.listing(job['user_id'],job_id=job['id']),[])
        self.assertEqual(json.loads(db.query_one('SELECT outcome_json FROM kilas_workspace_executions WHERE id=?',(run,))['outcome_json'])['large_storage_reserved'],large.STORAGE_RESERVATION)
        with jobs.transaction() as conn:
            work_artifacts.persist(conn,job,step,file);large.published(conn,job,run)
        artifact=work_artifacts.listing(job['user_id'],job_id=job['id'])[0]
        link=f'/kilas-ai/agent/jobs/{job["id"]}/artifacts/{artifact["id"]}'
        response=self.client(self.owner).get(link+'?download=1');self.assertEqual(response.data,raw)
        self.assertEqual(self.client(self.other).get(link).status_code,404)
        self.assertNotIn('large_storage_reserved',json.loads(db.query_one('SELECT outcome_json FROM kilas_workspace_executions WHERE id=?',(run,))['outcome_json']))
    def test_mock_remote_stream_lifecycle_cleanup_and_idempotency(self):
        from types import SimpleNamespace
        files={'large.pdf':LargeTests.files['large.pdf']};value=manifest(files)
        vm,bodies=sandbox(files);vm.sandbox_id='synthetic-large-vm'
        vm.get_info.return_value=SimpleNamespace(cpu_count=2,memory_mb=4096,network={'allow_public_traffic':False},allow_internet_access=False,lifecycle={'on_timeout':'kill','auto_resume':False})
        vm.commands.run.return_value=SimpleNamespace(stdout='ready')
        read=vm.files.read.side_effect
        report={'exit_code':0,'stdout':'42\n','stderr':'','outputs':value,'timed_out':False,'evidence':f.BOUNDARY}
        vm.files.read.side_effect=lambda path,**kw:f.Stream(json.dumps(report).encode()) if path.endswith('done.json') else read(path,**kw)
        create=self.enabled(vm);job,operation=self.claimed();step=self.step(job)
        with patch.object(large,'capability',return_value={'available':True}):
            outcome=remote.execute(job,operation,f.SOURCE)
            self.assertEqual(outcome['large_files'][0]['content'],files['large.pdf']);self.assertTrue(outcome['cleanup_confirmed']);vm.kill.assert_called_once()
            persisted=json.loads(db.query_one('SELECT outcome_json FROM kilas_workspace_executions')['outcome_json']);self.assertNotIn('large_files',persisted)
            self.assertEqual(persisted['large_storage_reserved'],large.STORAGE_RESERVATION)
            with self.assertRaisesRegex(ValueError,'already_dispatched'):remote.execute(job,operation,f.SOURCE)
            self.assertEqual(create.call_count,1)
        result=worker.result(outcome);self.assertTrue(runner.finish(job,job['lease_token'],step,result))
        self.assertEqual(len(work_artifacts.listing(self.owner,job_id=job['id'])),2)
    def test_interrupted_remote_transfer_kills_and_keeps_no_partial_artifacts(self):
        from types import SimpleNamespace
        files={'large.pdf':LargeTests.files['large.pdf']};vm,bodies=sandbox(files,True);vm.sandbox_id='synthetic-interrupt-vm'
        vm.get_info.return_value=SimpleNamespace(cpu_count=2,memory_mb=4096,network={'allow_public_traffic':False},allow_internet_access=False,lifecycle={'on_timeout':'kill','auto_resume':False})
        vm.commands.run.return_value=SimpleNamespace(stdout='ready');read=vm.files.read.side_effect
        report={'exit_code':0,'stdout':'','stderr':'','outputs':manifest(files),'timed_out':False,'evidence':f.BOUNDARY}
        vm.files.read.side_effect=lambda path,**kw:f.Stream(json.dumps(report).encode()) if path.endswith('done.json') else read(path,**kw)
        create=self.enabled(vm);job,operation=self.claimed()
        with patch.object(large,'capability',return_value={'available':True}):
            with self.assertRaisesRegex(ValueError,'remote_execution_error'):remote.execute(job,operation,f.SOURCE)
        vm.kill.assert_called_once();self.assertTrue(bodies[0].closed);self.assertEqual(create.call_count,1)
        row=db.query_one('SELECT state,cleanup_confirmed,outcome_json FROM kilas_workspace_executions')
        self.assertEqual(row['state'],'UNKNOWN');self.assertEqual(row['cleanup_confirmed'],1);self.assertEqual(json.loads(row['outcome_json'])['large_storage_reserved'],0)
        self.assertEqual(work_artifacts.listing(self.owner,job_id=job['id']),[])


if __name__=='__main__':unittest.main()


class ProofTests(unittest.TestCase):
    setUp=StorageTests.setUp
    enabled=StorageTests.enabled
    @classmethod
    def setUpClass(cls):
        import app
        cls.app=app.app
        from kilas_ai import workspace_large_fixture
        cls.files=workspace_large_fixture.documents()
    def remote(self,success=True):
        from types import SimpleNamespace
        from kilas_ai import workspace_large_fixture as fixture,workspace_acceptance_fixture as boundary
        files=self.files if success else {}
        vm,bodies=sandbox(files);vm.sandbox_id='synthetic-proof-vm'
        vm.get_info.return_value=SimpleNamespace(cpu_count=2,memory_mb=4096,network={'allow_public_traffic':False},allow_internet_access=False,lifecycle={'on_timeout':'kill','auto_resume':False})
        vm.commands.run.return_value=SimpleNamespace(stdout='ready')
        value={'packages':{k:True for k in fixture.PACKAGES},'totals':fixture.TOTALS,'documents':True,'sizes':fixture.SIZES,'boundary':{'answer':42,'checks':boundary.CHECKS,'evidence':{'data':{'written_bytes':66000000,'available_before':66000000,'requested_bytes':66060288,'errno':28},'inodes':{'created_files':4000,'available_before':4000,'errno':28}}}}
        report={'exit_code':0 if success else 1,'stdout':'KILAS_ACCEPTANCE_OK\nKILAS_LARGE_OK '+json.dumps(value)+'\n','stderr':'','outputs':manifest(files),'timed_out':False,'evidence':f.BOUNDARY,'collector_evidence':{'checks':['file_limit','job_limit','file_count'],'private_staging':True,'post_stop':True}}
        read=vm.files.read.side_effect
        vm.files.read.side_effect=lambda path,**kw:f.Stream(json.dumps(report).encode()) if path.endswith('done.json') else read(path,**kw)
        return vm,report
    def test_once_real_documents_download_handler_cleanup_gate_and_exact_fingerprint(self):
        from kilas_ai import workspace_large_proof as proof,usage
        vm,report=self.remote();create=self.enabled(vm)
        with patch.object(proof,'OWNER_ID',self.owner),patch.dict(os.environ,{'KILAS_LARGE_PROOF_APPROVAL':proof.POLICY,'KILAS_LARGE_OUTPUT_APPROVAL':''}):
            proof.run(self.app);proof.run(self.app)
            self.assertEqual(create.call_count,1);vm.kill.assert_called_once();self.assertTrue(proof.accepted())
            self.assertFalse(large.capability(self.owner)['available'])
            with patch.dict(os.environ,{'KILAS_LARGE_OUTPUT_APPROVAL':proof.POLICY}):self.assertTrue(large.capability(self.owner)['available'])
            with patch.object(large,'fingerprint',return_value='changed'):self.assertFalse(proof.accepted())
        row=db.query_one('SELECT outcome_json FROM kilas_workspace_executions');value=json.loads(row['outcome_json'])
        self.assertTrue(value['download_checked']);self.assertFalse(value['customer_metered']);self.assertNotIn('large_storage_reserved',value)
        self.assertEqual(db.query_one("SELECT COUNT(*) AS n FROM kilas_ai_usage WHERE user_id=? AND operation_type='CODE_EXECUTION'",(self.owner,))['n'],0)
    def test_failed_fixture_cannot_activate_or_retry(self):
        from kilas_ai import workspace_large_proof as proof
        vm,report=self.remote(False);create=self.enabled(vm)
        with patch.object(proof,'OWNER_ID',self.owner),patch.dict(os.environ,{'KILAS_LARGE_PROOF_APPROVAL':proof.POLICY,'KILAS_LARGE_OUTPUT_APPROVAL':proof.POLICY}):
            proof.run(self.app);proof.run(self.app)
            self.assertEqual(create.call_count,1);self.assertFalse(proof.accepted());self.assertFalse(large.capability(self.owner)['available'])
    def test_full_pass_predicate_rejects_each_missing_mandatory_check(self):
        from kilas_ai import workspace_large_proof as proof,workspace_large_fixture as fixture
        vm,report=self.remote();outcome={'exit_code':0,'timed_out':False,'customer_metered':False,'sandbox_created':True,'cleanup_confirmed':True,'elapsed_ms':1000,'stderr':'','stdout':report['stdout'],'boundary_evidence':f.BOUNDARY,'collector_evidence':report['collector_evidence'],'large_files':[{'filename':n,'mime_type':formats.MIMES[n.rsplit('.',1)[1]],'content':r} for n,r in self.files.items()],'host_memory_before':{'limit':536870912,'current':104857600,'peak':209715200},'host_memory_after':{'limit':536870912,'current':104857600,'peak':209715200}}
        self.assertTrue(proof.passed(outcome))
        for key,value in [('cleanup_confirmed',False),('customer_metered',True),('elapsed_ms',60000),('collector_evidence',{}),('large_files',[]),('host_memory_after',{'limit':536870912,'current':104857600,'peak':500000000})]:
            with self.subTest(key=key):self.assertFalse(proof.passed({**outcome,key:value}))
        value=json.loads(report['stdout'].split('KILAS_LARGE_OK ')[1]);value['packages']['reportlab']=False
        self.assertFalse(proof.passed({**outcome,'stdout':'KILAS_ACCEPTANCE_OK\nKILAS_LARGE_OK '+json.dumps(value)}))
    def test_actual_trusted_collector_above_limit_checks_locally(self):
        import shutil
        with tempfile.TemporaryDirectory() as folder,patch.object(guest,'UID',os.getuid(),create=True):
            original=Path
            def path(value):return original(folder)/'checks' if value=='/root/kilas-collector-check' else original(value)
            with patch.object(guest,'Path',side_effect=path):
                value=guest.collector_acceptance()
            self.assertEqual(value,{'checks':['file_limit','job_limit','file_count'],'private_staging':True,'post_stop':True})
            self.assertEqual(list(Path(folder).iterdir()),[])

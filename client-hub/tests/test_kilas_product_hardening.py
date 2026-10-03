"""Focused synthetic product boundaries; no external providers or production writes."""
import io
import os
import unittest
from datetime import datetime, timezone
from unittest.mock import patch
from werkzeug.datastructures import FileStorage
import test_kilas_ai_unified as base
from kilas_ai import attachments, work_attachments, work_schedule, unified_runtime, agent_store, agent_chat, providers


def upload(raw,name,mime):
    return FileStorage(stream=io.BytesIO(raw),filename=name,content_type=mime)


class HardeningTests(base.UnifiedTests):
    def test_root_ignores_stale_assist_selection(self):
        client=self.client()
        with client.session_transaction() as session:session['active_product']='brain'
        response=client.get('/')
        self.assertTrue(response.location.endswith('/products/start'))
        with client.session_transaction() as session:self.assertEqual(session['active_product'],'kilas_ai')

    def test_connections_not_a_customer_surface(self):
        client=self.client()
        response=client.get('/kilas-ai/agent?view=connections')
        self.assertEqual(response.status_code,303)
        html=client.get('/kilas-ai/agent').get_data(as_text=True)
        self.assertNotIn('>Connections</a>',html)
        self.assertNotIn('Hubungkan Google',html)

    def test_future_news_never_searches_now(self):
        from kilas_ai import autonomous_store
        with patch.object(unified_runtime.tools,'web_search_steps',side_effect=AssertionError('Search before wake')):
            self.submit('buatin gw berita terbaru tentang amerika bsk jam 5 pagi ya')
        jobs=autonomous_store.list_jobs(self.owner)
        self.assertEqual(len(jobs),1)
        self.assertEqual(jobs[0]['mode'],'SCHEDULED')
        from kilas_ai import autonomous_planner
        plan=autonomous_planner.propose(jobs[0],[])
        self.assertEqual(plan['steps'][0]['worker'],'WEB')
        self.assertEqual(len(agent_store.messages(self.owner,conversation_id=self.conversation)),2)

    def test_local_schedule_normalization(self):
        now=datetime(2026,10,3,12,tzinfo=timezone.utc)
        for zone,hour in [('Asia/Jakarta',22),('Asia/Makassar',21),('Asia/Jayapura',20)]:
            spec=work_schedule.parse('buatin berita amerika bsk pukul 5 pagi',zone,now)
            self.assertEqual(spec['next_run_at'],datetime(2026,10,3,hour,tzinfo=timezone.utc))
        self.assertEqual(unified_runtime.intent('cari berita terbaru lusa jam 8'),'WORK')

    def test_daily_news_is_search_work_not_a_reminder(self):
        import json
        from kilas_ai import autonomous_planner
        self.submit('tiap pagi jam 7 cari berita ekonomi terbaru')
        job=self.jobs.list_jobs(self.owner)[0]
        self.assertEqual(job['mode'],'RECURRING')
        self.assertNotIn('reminder',json.loads(job['checkpoint_json']))
        self.assertEqual(autonomous_planner.propose(job,[])['steps'][0]['worker'],'WEB')

    def test_visual_followup_uses_concept(self):
        for prompt in ('iya dalam gambar bisa?','sekarang bikin gambarnya','buat versi visualnya'):
            self.assertEqual(unified_runtime.intent(prompt,previous_answer='Konsep poster kopi dengan ilustrasi cangkir.'),'IMAGE_GENERATE')

    def test_visual_followup_reaches_real_image_worker(self):
        from kilas_ai import autonomous_planner
        for prompt in ('iya dalam gambar bisa?','weh iya gambar aja','buat versi visualnya'):
            self.conversation=agent_store.new_conversation(self.owner)
            agent_store.append(self.owner,'assistant','Konsep poster kopi dengan ilustrasi cangkir dan warna krem.',self.conversation)
            self.submit(prompt)
            jobs=self.jobs.list_jobs(self.owner,conversation_id=self.conversation)
            self.assertEqual(len(jobs),1,prompt)
            plan=autonomous_planner.propose(jobs[0],[])
            self.assertEqual(plan['steps'][0]['worker'],'IMAGE')
            self.assertIn('cangkir',jobs[0]['instruction'])

    def test_spreadsheet_formulas_are_data_not_executed(self):
        from openpyxl import Workbook
        book=Workbook();book.active.title='Anggaran';book.active.append(['Jumlah',10]);book.active.append(['Total','=SUM(B1:B1)'])
        raw=io.BytesIO();book.save(raw)
        item=work_attachments.prepare_many([upload(raw.getvalue(),'budget.xlsx',work_attachments.work_artifacts.MIMES['xlsx'])],'FREE')[0]
        self.assertIn('Sheet: Anggaran',item['extracted_text'])
        self.assertIn('=SUM(B1:B1)',item['extracted_text'])
        self.assertIn('not recalculated',item['extracted_text'])

    def test_presentation_preserves_slide_boundaries(self):
        from pptx import Presentation
        deck=Presentation();slide=deck.slides.add_slide(deck.slide_layouts[1]);slide.shapes.title.text='Strategi';slide.placeholders[1].text='Pertumbuhan';slide.notes_slide.notes_text_frame.text='Catatan QA'
        raw=io.BytesIO();deck.save(raw)
        item=work_attachments.prepare_many([upload(raw.getvalue(),'strategy.pptx',work_attachments.work_artifacts.MIMES['pptx'])],'FREE')[0]
        self.assertIn('Slide 1',item['extracted_text']);self.assertIn('Strategi',item['extracted_text']);self.assertIn('Catatan QA',item['extracted_text'])

    def test_scanned_pdf_uses_bounded_vision(self):
        from PIL import Image
        image=Image.new('RGB',(300,300),'white');raw=io.BytesIO();image.save(raw,'PDF')
        with patch.dict(os.environ,{'OPENAI_API_KEY':'synthetic-test-key'}):
            item=attachments.prepare(upload(raw.getvalue(),'scan.pdf','application/pdf'))
        self.assertEqual(len(item['vision_pages']),1)
        blocks=attachments.prompt_content('baca scan ini',[item])
        self.assertEqual(blocks[1]['type'],'image_url')
        self.assertIn('halaman pertama',blocks[0]['text'])

    def test_utf16_and_structured_csv(self):
        item=attachments.prepare(upload('Nama;Jumlah\nKopi;30'.encode('utf-16'),'sample.csv','text/csv'))
        self.assertIn('Row 1: Nama | Jumlah',item['extracted_text'])
        self.assertIn('Row 2: Kopi | 30',item['extracted_text'])

    def test_uploaded_text_survives_followup(self):
        client=self.client()
        with patch.object(providers,'stream',side_effect=lambda *a,**k:base.answer('Anggaran tercatat.')):
            response=client.post('/kilas-ai/agent/chat',data={'csrf_token':'work-csrf','message':'ringkas ini','operation_key':'hardening-upload-key-12345','conversation_id':str(self.conversation),'source_files':(io.BytesIO(b'Anggaran keuangan 30 juta'),'facts.txt')},content_type='multipart/form-data')
        self.assertNotEqual(response.status_code,400)
        with patch.object(providers,'stream',side_effect=lambda *a,**k:base.answer('Bagian keuangan mencatat anggaran 30 juta.')) as called:
            self.submit('bagian keuangannya gimana?')
        self.assertIn('Anggaran keuangan 30 juta',str(called.call_args.args[1]))

    def test_source_context_owner_and_conversation_isolation(self):
        from kilas_ai import agent_attachments
        agent_store.append(self.owner,'user','file',self.conversation,attachments=[{
            'filename':'private.txt','mime_type':'text/plain','byte_size':10,'content':b'owner-only','extracted_text':'owner-only'}])
        self.assertEqual(agent_attachments.sources(self.other,self.conversation),[])
        other_chat=agent_store.new_conversation(self.owner)
        self.assertEqual(agent_attachments.sources(self.owner,other_chat),[])
        self.assertEqual(agent_attachments.sources(self.owner,self.conversation)[0]['text'],'owner-only')

    def test_multiple_long_files_keep_all_source_boundaries(self):
        files=[{'filename':f'proposal-{n}.txt','mime_type':'text/plain','extracted_text':str(n)*12000} for n in range(5)]
        prompt=attachments.prompt_content('bandingkan semua',files)
        for n in range(5):self.assertIn(f'proposal-{n}.txt',prompt)
        self.assertEqual(prompt.count('</isi_lampiran>'),5)
        self.assertLessEqual(len(prompt),16000)

    def test_uploaded_document_analysis_does_not_create_another_file(self):
        prepared=[{'filename':'qa.pdf','mime_type':'application/pdf','extracted_text':'Document facts'}]
        self.assertEqual(unified_runtime.intent('ringkas PDF ini',prepared),'CHAT')
        self.assertEqual(unified_runtime.intent('ringkas PDF ini jadiin PDF baru',prepared),'DOCUMENT')


if __name__=='__main__':unittest.main()

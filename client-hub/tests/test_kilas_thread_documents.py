"""Real thread document rendering/downloads with mocked authors; no provider/VM IO."""
import io,json,random,unittest
from unittest.mock import patch
from PIL import Image
from pypdf import PdfReader
from pptx import Presentation
import test_kilas_chat_models as fixture
import db
from kilas_ai import providers,store,usage,routing,unified_runtime,workspace_projects

PDF='# Laporan Sintetis\n\n## Ringkasan\n\nPendapatan sintetis 410000, biaya 226000, laba 184000.'
PPTX='# Presentasi Sintetis\n\n## Ringkasan\n- Pendapatan 410000\n- Biaya 226000\n- Laba 184000'


class ThreadDocuments(unittest.TestCase):
    setUp=fixture.RouteTests.setUp
    client=fixture.RouteTests.client
    submit=fixture.RouteTests.submit

    def stream(self,source,finish='stop'):
        def generate(mode,context,**kwargs):
            model=context.selected_model
            yield {'type':'provider','provider':'openai','model':model}
            yield {'type':'model_metadata','provider':'openai','requested_model':model,'returned_model':model}
            yield {'type':'delta','text':source}
            yield {'type':'usage','input_tokens':100,'output_tokens':150}
            yield {'type':'finish','reason':finish}
        return generate

    def test_small_pdf_pptx_exact_selected_model_download_reload_and_idempotency(self):
        for format,source,model in [('pdf',PDF,'gpt-6-luna'),('pptx',PPTX,'gpt-6.1-sol')]:
            key='thread_document_'+format+'_00000001'
            with patch.object(providers,'stream',side_effect=self.stream(source)) as author,patch('e2b.Sandbox.create',side_effect=AssertionError('No VM for document rendering')):
                response=self.submit(model,'Create a '+format+' report with synthetic data.',agent=False,key=key)
                self.assertIn('event: file',response.text);self.assertEqual(author.call_args.args[1].selected_model,model)
                self.submit(model,'Create a '+format+' report with synthetic data.',agent=False,key=key)
                author.assert_called_once()
            file=store.attachment_list(self.owner,self.thread)[-1]
            path=f'/kilas-ai/threads/{self.thread}/attachments/{file["id"]}'
            downloaded=self.client().get(path);self.assertEqual(downloaded.status_code,200)
            self.assertEqual(downloaded.data,bytes(store.attachment(self.owner,self.thread,file['id'])['content']))
            self.assertEqual(self.client(self.other).get(path).status_code,404)
            another=store.create_thread(self.owner)
            self.assertEqual(self.client().get(f'/kilas-ai/threads/{another}/attachments/{file["id"]}').status_code,404)
            self.assertIn('attachment;',downloaded.headers['Content-Disposition'])
            self.assertIn(file['filename'],self.client().get(f'/kilas-ai/threads/{self.thread}').text)
            if format=='pdf':self.assertIn('184000',PdfReader(io.BytesIO(downloaded.data),strict=True).pages[0].extract_text())
            else:
                deck=Presentation(io.BytesIO(downloaded.data));self.assertEqual(len(deck.slides),2)
                self.assertIn('184000',' '.join(s.text for slide in deck.slides for s in slide.shapes if s.has_text_frame))
            row=db.query_one('SELECT operation_type,status,model FROM kilas_ai_usage WHERE user_id=? AND operation_key=?',(self.owner,key))
            self.assertEqual((row['operation_type'],row['status'],row['model']),('PDF' if format=='pdf' else 'CHAT','COMPLETE',model))

    def test_representative_large_pdf_and_multilingual_pptx_through_actual_thread(self):
        rng=random.Random(19)
        source='# Laporan Sintetis Besar\n\n'+'\n\n'.join(f'## Bagian {i+1}\n\n'+('Data sintetis untuk peninjauan: pendapatan 410000 biaya 226000 laba 184000. '*16) for i in range(12))
        png=io.BytesIO();Image.frombytes('RGB',(640,480),rng.randbytes(640*480*3)).save(png,format='PNG')
        with patch.object(providers,'stream',side_effect=self.stream(source)):
            result=self.submit('gpt-6-luna','Buat PDF laporan besar sintetis dengan logo terlampir.',agent=False,attachments=(io.BytesIO(png.getvalue()),'logo.png'))
        self.assertIn('event: file',result.text)
        file=store.attachment_list(self.owner,self.thread)[-1]
        raw=self.client().get(f'/kilas-ai/threads/{self.thread}/attachments/{file["id"]}').data
        pages=PdfReader(io.BytesIO(raw),strict=True).pages;self.assertGreaterEqual(len(pages),7);self.assertGreater(len(raw),65536)
        self.assertIn('Bagian 12',' '.join(p.extract_text() for p in pages))
        source='# Presentasi Sintetis Besar\n\n'+'\n\n'.join(f'## Bagian {i+1}\n'+'\n'.join('- '+''.join(chr(rng.randrange(0x4e00,0x9fff)) for _ in range(145)) for _ in range(7)) for i in range(15))
        with patch.object(providers,'stream',side_effect=self.stream(source)):
            result=self.submit('gpt-6.1-sol','Buat PPTX besar dari contoh sintetis multibahasa.',agent=False)
        self.assertIn('event: file',result.text)
        file=store.attachment_list(self.owner,self.thread)[-1]
        raw=self.client().get(f'/kilas-ai/threads/{self.thread}/attachments/{file["id"]}').data
        deck=Presentation(io.BytesIO(raw));self.assertEqual(len(deck.slides),16);self.assertGreater(len(raw),65536)
        text='\n'.join(s.text for slide in deck.slides for s in slide.shapes if s.has_text_frame)
        for line in source.splitlines():
            if line:self.assertIn(line.lstrip('-# '),text)
        self.assertEqual(self.client(self.other).get(f'/kilas-ai/threads/{self.thread}/attachments/{file["id"]}').status_code,404)
        print(f'THREAD_DOCUMENT_EVIDENCE pdf_bytes={len(bytes(store.attachment(self.owner,self.thread,store.attachment_list(self.owner,self.thread)[1]["id"])["content"]))} pdf_pages={len(pages)} pptx_bytes={len(raw)} pptx_slides={len(deck.slides)}')

    def test_pptx_revision_uses_only_owned_presentation_source(self):
        with patch.object(providers,'stream',side_effect=self.stream(PPTX)):
            self.submit('gpt-6-luna','Buat presentasi PPTX sintetis.',agent=False)
        with patch.object(providers,'stream',side_effect=self.stream(PPTX+'\n- Revisi selesai')) as author:
            self.submit('gpt-6-luna','Revisi PPTX tadi dengan satu poin lagi.',agent=False)
        self.assertIn('Pendapatan 410000',str(author.call_args.args[1]));self.assertEqual(len(store.attachment_list(self.owner,self.thread)),2)

    def test_incomplete_invalid_and_storage_failure_never_publish_success_or_raw_source(self):
        for index,(source,finish) in enumerate([(PPTX,'length'),('print("not a presentation")','stop')]):
            with patch.object(providers,'stream',side_effect=self.stream(source,finish)):
                result=self.submit('gpt-6-luna','Buat PPTX sintetis.',agent=False,key='invalid_document_'+str(index).zfill(18))
            self.assertIn('event: error',result.text);self.assertNotIn('event: file',result.text)
            self.assertFalse(store.attachment_list(self.owner,self.thread))
        original=usage._query
        def fail_file(conn,sql,*args,**kwargs):
            if sql.startswith('INSERT INTO kilas_ai_attachments'):raise ValueError('synthetic storage failure')
            return original(conn,sql,*args,**kwargs)
        with patch.object(providers,'stream',side_effect=self.stream(PPTX)),patch.object(usage,'_query',side_effect=fail_file):
            result=self.submit('gpt-6-luna','Buat PPTX sintetis.',agent=False,key='storage_document_00000001')
        self.assertIn('event: error',result.text);self.assertFalse(store.attachment_list(self.owner,self.thread))
        self.assertIsNone(store.operation(self.owner,self.thread,'storage_document_00000001:assistant'))
        self.assertFalse(db.query_all("SELECT id FROM kilas_ai_messages WHERE thread_id=? AND role='assistant'",(self.thread,)))

    def test_quota_auth_and_csrf_stop_before_authoring(self):
        with patch.object(providers,'stream') as author:
            with patch.object(usage,'reserve',side_effect=usage.UsageLimit('Synthetic quota exhausted')):
                self.assertEqual(self.submit('gpt-6-luna','Buat PPTX sintetis.',agent=False).status_code,429)
            path=f'/kilas-ai/threads/{self.thread}/send'
            self.assertEqual(self.client(self.other).post(path,data={'content':'Buat PPTX','operation_key':'foreign_document_00000001','csrf_token':'work-csrf'}).status_code,404)
            self.assertEqual(self.client().post(path,json={'content':'Buat PPTX','operation_key':'csrf_document_00000001'}).status_code,400)
            author.assert_not_called()

    def test_document_routing_preserves_questions_negation_and_explicit_execution(self):
        for text in ('Create PDF report','Generate PDF report','Buat PDF besar'):
            self.assertEqual(routing.tool_for(text),'PDF')
        for text in ('Buat PPTX','Buat presentasi','Create a PowerPoint presentation'):
            self.assertEqual(routing.tool_for(text),'PPTX');self.assertEqual(unified_runtime.intent(text),'DOCUMENT')
            from kilas_ai import work_documents
            self.assertEqual(work_documents.intent(text),'pptx')
        for text in ('Apa itu PPTX?',"Don't create PPTX; explain it.",'Jangan buat PDF baru.', 'Jalankan Python untuk buat PPTX', 'Create a Python project for PPTX'):
            self.assertEqual(routing.tool_for(text),'CHAT')
        for text in ("Don't create PPTX; explain it.",'Jangan buat PDF baru.'):
            self.assertEqual(unified_runtime.intent(text),'CHAT')

    def test_authoring_limits_match_each_transport_without_enabling_large(self):
        from kilas_ai import workspace_large_outputs as large
        small=workspace_projects.export_instruction();big=workspace_projects.export_instruction(True)
        for text in ('65536','180000','45000'):self.assertIn(text,small);self.assertNotIn(text,big)
        for limit in (large.MAX_FILE,large.MAX_JOB):self.assertIn(str(limit),big)
        self.assertIn('at most 4 files',big);self.assertFalse(large.capability(self.owner)['available'])


if __name__=='__main__':unittest.main()

"""Generated local fixtures/mock execution; no provider or real VM calls."""
import ast,base64,copy,hashlib,io,json,os,tempfile,unittest,warnings,zipfile
from pathlib import Path
from unittest.mock import patch
from kilas_ai import workspace_output_formats as formats,workspace_artifact_transport as transport
from kilas_ai import workspace_execution_worker as worker,workspace_artifact_exporter as exporter
from kilas_ai import work_artifacts,workspace_execution as execution,workspace_execution_proof as proof,workspace_execution_rollout as rollout
import test_kilas_workspace_execution as f
import db


def fixtures():
    from reportlab.pdfgen import canvas
    pdf=io.BytesIO();c=canvas.Canvas(pdf);c.drawString(50,800,'Synthetic sales: revenue 410000 cost 226000 profit 184000');c.save()
    from docx import Document
    doc=Document();doc.add_heading('Synthetic sales',0);doc.add_paragraph('Revenue 410000, cost 226000, profit 184000.');docx=io.BytesIO();doc.save(docx)
    from openpyxl import Workbook
    from openpyxl.chart import BarChart,Reference
    book=Workbook();sheet=book.active
    for row in [('Item','Units','Price','Unit cost','Revenue','Cost','Profit'),('Kopi',10,20000,12000,'=B2*C2','=B2*D2','=E2-F2'),('Teh',8,15000,7000,'=B3*C3','=B3*D3','=E3-F3'),('Roti',5,18000,10000,'=B4*C4','=B4*D4','=E4-F4')]:sheet.append(row)
    chart=BarChart();chart.add_data(Reference(sheet,min_col=5,max_col=7,min_row=1,max_row=4),titles_from_data=True);sheet.add_chart(chart,'I2');xlsx=io.BytesIO();book.save(xlsx)
    from pptx import Presentation
    deck=Presentation();slide=deck.slides.add_slide(deck.slide_layouts[5]);slide.shapes.title.text='Synthetic sales: profit 184000';pptx=io.BytesIO();deck.save(pptx)
    from PIL import Image
    images={}
    for suffix,kind in [('png','PNG'),('jpg','JPEG'),('webp','WEBP'),('gif','GIF'),('bmp','BMP')]:
        out=io.BytesIO();Image.new('RGB',(16,16),'blue').save(out,format=kind);images['chart.'+suffix]=out.getvalue()
    return {'report.pdf':pdf.getvalue(),'report.docx':docx.getvalue(),'sales.xlsx':xlsx.getvalue(),'slides.pptx':pptx.getvalue(),**images,'sales.csv':b'item,revenue,cost,profit\nTotal,410000,226000,184000\n','totals.json':b'{"revenue":410000,"cost":226000,"profit":184000}'}


def frames(files):
    result={};number=0
    for name,raw in files.items():
        count=(len(raw)+44999)//45000
        for part in range(count):
            result[f'kilas-export-{number:03}.json']=json.dumps({'protocol':'kilas.artifact.v1','name':name,'part':part,'parts':count,'sha256':hashlib.sha256(raw).hexdigest(),'base64':base64.b64encode(raw[part*45000:(part+1)*45000]).decode()});number+=1
    result['kilas-runtime.json']=json.dumps({'protocol':'kilas.runtime.v1','python':'3.14.3','packages':{k:False for k in ('reportlab','docx','openpyxl','xlsxwriter','pptx','PIL')}})
    return result


def outcome(outputs):
    return {'execution':'completed','provider':'e2b','run_id':1,'exit_code':0,'timed_out':False,'stdout':'410000 226000 184000\n','stderr':'','outputs':outputs,'cleanup_confirmed':True,'customer_metered':True}


class FormatTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):cls.files=fixtures()
    def test_real_formats_roundtrip_and_editable_office_formula_chart(self):
        for name,raw in self.files.items():
            with self.subTest(name=name):
                mime=formats.MIMES[formats.filename(name)];self.assertEqual(formats.validate(name,mime,raw),raw)
                decoded,runtime=transport.collect(frames({name:raw}));self.assertEqual(decoded[0]['content'],raw)
                bundle=transport.bundle(decoded);self.assertIsNotNone(bundle)
                with zipfile.ZipFile(io.BytesIO(bundle['content'])) as archive:self.assertEqual(archive.read(name),raw)
        from docx import Document
        self.assertEqual(Document(io.BytesIO(self.files['report.docx'])).paragraphs[0].text,'Synthetic sales')
        from openpyxl import load_workbook
        book=load_workbook(io.BytesIO(self.files['sales.xlsx']));self.assertEqual(book.active['G2'].value,'=E2-F2');self.assertEqual(len(book.active._charts),1)
        from pptx import Presentation
        self.assertEqual(len(Presentation(io.BytesIO(self.files['slides.pptx'])).slides),1)
    def test_magic_mime_extension_executable_and_svg_raster_denied(self):
        for name,mime,raw in [('fake.pdf','application/pdf',self.files['chart.png']),('fake.xlsx',formats.MIMES['xlsx'],self.files['report.docx']),('fake.png','image/png',b'<svg onload="x"/>'),('program.exe','application/octet-stream',b'MZbad'),('fake.txt','text/plain',b'\x7fELFbad')]:
            with self.subTest(name=name),self.assertRaises(Exception):formats.validate(name,mime,raw)
    def test_pdf_active_javascript_and_embedded_launch_rejected(self):
        from pypdf import PdfReader,PdfWriter
        for mode in ('javascript','attachment'):
            writer=PdfWriter();writer.append(PdfReader(io.BytesIO(self.files['report.pdf'])))
            if mode=='javascript':writer.add_js('app.alert(1)')
            else:writer.add_attachment('code.txt',b'payload')
            raw=io.BytesIO();writer.write(raw)
            with self.subTest(mode=mode),self.assertRaisesRegex(ValueError,'active_pdf'):formats.validate('report.pdf','application/pdf',raw.getvalue())
    def test_office_external_macro_and_formula_denied(self):
        raw=self.files['report.docx']
        for name,body in [('_rels/.rels',b'<Relationships><Relationship TargetMode="External" Target="https://bad"/></Relationships>'),('word/vbaProject.bin',b'macro')]:
            existing=formats.entries(raw);existing[name]=body;out=io.BytesIO()
            with zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED) as archive:
                for path,data in existing.items():archive.writestr(path,data)
            with self.subTest(name=name),self.assertRaises(ValueError):formats.validate('report.docx',formats.MIMES['docx'],out.getvalue())
        from openpyxl import Workbook
        for formula in ('=WEBSERVICE("https://bad")','=HYPERLINK("https://bad")',"=cmd|' /C calc'!A0",'=[outside.xlsx]Sheet1!A1'):
            book=Workbook();book.active['A1']=formula;out=io.BytesIO();book.save(out)
            with self.subTest(formula=formula),self.assertRaisesRegex(ValueError,'unsafe_formula'):formats.validate('sales.xlsx',formats.MIMES['xlsx'],out.getvalue())
    def test_archive_traversal_duplicates_encryption_bomb_nesting(self):
        for names in [('../escape.txt',),('/absolute.txt',),('a.txt','A.txt'),('.env',),('huge.txt',)]:
            out=io.BytesIO()
            with zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED) as archive:
                for name in names:archive.writestr(name,b'x'*100000 if name=='huge.txt' else b'x')
            with self.subTest(names=names),self.assertRaises(ValueError):formats.validate('project.zip','application/zip',out.getvalue())
        inner=io.BytesIO()
        with zipfile.ZipFile(inner,'w') as archive:archive.writestr('safe.txt',b'safe')
        outer=io.BytesIO()
        with zipfile.ZipFile(outer,'w') as archive:archive.writestr('nested.zip',inner.getvalue())
        with self.assertRaisesRegex(ValueError,'nested_archive'):formats.validate('project.zip','application/zip',outer.getvalue())
        encrypted=bytearray(inner.getvalue());encrypted[6]|=1
        with self.assertRaises(ValueError):formats.validate('project.zip','application/zip',bytes(encrypted))
    def test_transport_tamper_truncation_alias_sequence_path_limit_base64(self):
        valid=frames({'report.pdf':self.files['report.pdf']})
        for change in ({'name':'../escape.pdf'},{'protocol':'v0'},{'part':True},{'parts':2},{'sha256':'0'*64},{'base64':'%%%bad'},{'extra':1},{'name':'execution-result.json'}):
            invalid=copy.deepcopy(valid);frame=json.loads(invalid['kilas-export-000.json']);frame.update(change);invalid['kilas-export-000.json']=json.dumps(frame)
            with self.subTest(change=change),self.assertRaises((ValueError,KeyError,TypeError)):transport.collect(invalid)
        invalid=copy.deepcopy(valid);invalid['kilas-export-001.json']=invalid['kilas-export-000.json']
        with self.assertRaisesRegex(ValueError,'collision'):transport.collect(invalid)
        invalid=copy.deepcopy(valid);invalid['kilas-export-007.json']=invalid.pop('kilas-export-000.json')
        with self.assertRaisesRegex(ValueError,'sequence'):transport.collect(invalid)
        with self.assertRaises(ValueError):transport.collect(frames({'a.txt':b'a'*65537}))
        with self.assertRaises(ValueError):transport.collect(frames({f'{n}.txt':b'a' for n in range(5)}))
        invalid=copy.deepcopy(valid);invalid['kilas-export-000.json']='{"protocol":"v0",'+invalid['kilas-export-000.json'][1:]
        with self.assertRaisesRegex(ValueError,'duplicate_json'):transport.collect(invalid)
    def test_fixed_exporter_automatically_frames_existing_binary_without_host_program_execution(self):
        old=Path.cwd()
        with tempfile.TemporaryDirectory() as folder:
            try:
                os.chdir(folder);Path('output').mkdir();Path('output/report.pdf').write_bytes(self.files['report.pdf']);exporter.export_outputs()
                raw={p.name:p.read_text() for p in Path('output').iterdir()};files,runtime=transport.collect(raw)
                self.assertEqual(files[0]['content'],self.files['report.pdf']);self.assertIn('unprivileged',runtime['evidence'])
                self.assertFalse(Path('output/report.pdf').exists())
            finally:os.chdir(old)
        files,entry=transport.prepare({'main.py':'print(42)'},'main.py');self.assertEqual(entry,'_kilas_entrypoint.py');ast.parse(files[entry]);self.assertIn('finally:',files[entry]);self.assertIn('runpy.run_path',files[entry])
        with self.assertRaises(ValueError):transport.prepare({'_kilas_artifacts.py':'bad'},'main.py')
        with self.assertRaises(ValueError):transport.prepare({'main.py':'print(42)'},'../private.py')
    def test_proven_bundle_and_finite_pin_unchanged(self):
        self.assertEqual(proof.fingerprint(),'84f9744af2820215762ee4d60bcc9db560ef5a2ed1176cc8448fe2999e3dbd7b')
        self.assertEqual(hashlib.sha256(Path(rollout.__file__).read_bytes()).hexdigest(),'1403ac10aba25b6fa29916991add1f32e143d6a8cb59b25153e0ef70e4336d83')
        self.assertTrue(rollout.accepts({'proof_passed':True,'proof_fingerprint':rollout.ORIGINAL_PROOF_FINGERPRINT}))


class ArtifactStorageTests(unittest.TestCase):
    def setUp(self):
        import uuid
        identity=self.id();self.id=lambda:identity+'-'+uuid.uuid4().hex[:8]
        f.ExecutionTests.setUp(self)
    claimed=f.ExecutionTests.claimed
    client=f.ExecutionTests.client
    @classmethod
    def setUpClass(cls):cls.files=fixtures()
    def test_mock_execution_real_storage_download_preview_and_cross_owner_denial(self):
        from kilas_ai import autonomous_runner as runner,autonomous_store as jobs
        job,step=self.claimed();step={'id':None,'worker':'EXECUTION'}
        # Real step/lease and existing transactional publication, no SDK.
        db.execute("INSERT INTO kilas_agent_steps(job_id,sequence,worker,action,instruction,input_json,idempotency_key,status) VALUES (?,1,'EXECUTION','python','Synthetic','{}',?,'RUNNING')",(job['id'],'binary-'+str(job['id'])))
        step=dict(db.query_one('SELECT * FROM kilas_agent_steps WHERE job_id=?',(job['id'],)))
        result=worker.result(outcome(frames({'report.pdf':self.files['report.pdf'],'chart.png':self.files['chart.png']})))
        with patch('e2b.Sandbox.create') as sdk:self.assertTrue(runner.finish(job,job['lease_token'],step,result));sdk.assert_not_called()
        files=work_artifacts.listing(self.owner,job_id=job['id']);self.assertEqual(len(files),4)
        for artifact in files:
            link=f'/kilas-ai/agent/jobs/{job["id"]}/artifacts/{artifact["id"]}'
            response=self.client().get(link+'?download=1');self.assertEqual(response.status_code,200);self.assertIn('attachment',response.headers['Content-Disposition']);self.assertEqual(response.headers['X-Content-Type-Options'],'nosniff')
            if artifact['name']=='report.pdf':self.assertEqual(response.data,self.files['report.pdf'])
            self.assertEqual(self.client(self.other).get(link).status_code,404)
        self.assertFalse(runner.finish(job,job['lease_token'],step,result));self.assertEqual(len(work_artifacts.listing(self.owner,job_id=job['id'])),4)
    def test_invalid_second_artifact_rolls_back_every_file(self):
        from kilas_ai import autonomous_runner as runner,agent_workers
        job,_=self.claimed();db.execute("INSERT INTO kilas_agent_steps(job_id,sequence,worker,action,instruction,input_json,idempotency_key,status) VALUES (?,1,'EXECUTION','python','Synthetic','{}',?,'RUNNING')",(job['id'],'atomic-'+str(job['id'])))
        step=dict(db.query_one('SELECT * FROM kilas_agent_steps WHERE job_id=?',(job['id'],)))
        result=agent_workers.Result('SUCCEEDED','Done',{},[{'binary_file':{'filename':'ok.txt','mime_type':'text/plain','content':b'ok'}},{'binary_file':{'filename':'bad.pdf','mime_type':'application/pdf','content':b'not PDF'}}],verified=True)
        with self.assertRaises(ValueError):runner.finish(job,job['lease_token'],step,result)
        self.assertEqual(work_artifacts.listing(self.owner,job_id=job['id']),[])
